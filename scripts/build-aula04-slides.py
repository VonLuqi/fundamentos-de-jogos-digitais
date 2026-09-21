#!/usr/bin/env python3
"""Gera os slides da Aula 04 a partir do template visual da Aula 01.

Saídas:
  assets/docs/aulas/aula04_plataformas_restricoes_slides.pptx
  assets/docs/aulas/aula04_plataformas_restricoes_slides.pdf  (via PowerPoint COM, se disponível)
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pptx.oxml import parse_xml

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "assets" / "docs" / "aulas" / "aula01_godot_slides.pptx"
OUT_PPTX = ROOT / "assets" / "docs" / "aulas" / "aula04_plataformas_restricoes_slides.pptx"
OUT_PDF = ROOT / "assets" / "docs" / "aulas" / "aula04_plataformas_restricoes_slides.pdf"

BG = RGBColor(0x0D, 0x0A, 0x10)
PANEL = RGBColor(0x18, 0x12, 0x1A)
GOLD = RGBColor(0xCF, 0xA7, 0x59)
GOLD_BRIGHT = RGBColor(0xF2, 0xD5, 0x9A)
GOLD_DARK = RGBColor(0x7A, 0x5C, 0x2E)
TEXT = RGBColor(0xEC, 0xE1, 0xD1)
TEXT_DIM = RGBColor(0xAB, 0x9C, 0x8A)
BLOOD = RGBColor(0x7A, 0x1F, 0x2B)

TITLE_FONT = "Cinzel"
BODY_FONT = "Georgia"
FOOTER = "AULA 04 | MÓDULO 1"


def _set_run(run, *, size_pt, bold=False, color=TEXT, font=BODY_FONT):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "cs", "ea"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = parse_xml(
                f'<a:{tag} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'
            )
            rPr.append(el)
        else:
            el.set("typeface", font)


def _clear_paragraphs(tf):
    p = tf.paragraphs[0]
    p.clear()
    for extra in list(tf.paragraphs)[1:]:
        p_elem = extra._p
        p_elem.getparent().remove(p_elem)
    return tf.paragraphs[0]


def add_textbox(slide, left, top, width, height, *, word_wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    return box, tf


def write_lines(tf, lines, *, size_pt=18, bold=False, color=TEXT, font=BODY_FONT, align=PP_ALIGN.LEFT, space_after=6):
    first = True
    for item in lines:
        if isinstance(item, tuple):
            text, opts = item[0], item[1]
        else:
            text, opts = item, {}
        p = _clear_paragraphs(tf) if first else tf.add_paragraph()
        first = False
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", space_after))
        run = p.add_run()
        run.text = text
        _set_run(
            run,
            size_pt=opts.get("size", size_pt),
            bold=opts.get("bold", bold),
            color=opts.get("color", color),
            font=opts.get("font", font),
        )


def paint_background(slide, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD_DARK
    bar.line.fill.background()
    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), height - Inches(0.08), width, Inches(0.08)
    )
    foot.fill.solid()
    foot.fill.fore_color.rgb = GOLD_DARK
    foot.line.fill.background()


def add_footer(slide, width, height, text=FOOTER):
    _, tf = add_textbox(
        slide,
        Inches(0.5),
        height - Inches(0.45),
        width - Inches(1.0),
        Inches(0.3),
    )
    write_lines(tf, [text], size_pt=10, color=TEXT_DIM, font=TITLE_FONT, align=PP_ALIGN.LEFT)


def add_num_badge(slide, number: str):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(0.28), Inches(0.7), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLOOD
    badge.line.color.rgb = GOLD
    badge.line.width = Pt(1)
    tf = badge.text_frame
    tf.word_wrap = False
    p = _clear_paragraphs(tf)
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = number
    _set_run(run, size_pt=14, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_title(slide, title: str, left=Inches(1.35), top=Inches(0.28), width=Inches(8.2)):
    _, tf = add_textbox(slide, left, top, width, Inches(0.55))
    write_lines(tf, [title], size_pt=22, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_panel(slide, left, top, width, height):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = GOLD_DARK
    panel.line.width = Pt(1.25)
    return panel


def delete_all_slides(prs: Presentation):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def new_slide(prs: Presentation):
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    for shape in list(slide.shapes):
        sp = shape._element
        sp.getparent().remove(sp)
    paint_background(slide, prs.slide_width, prs.slide_height)
    return slide


def build_deck() -> Presentation:
    if not SRC.exists():
        raise SystemExit(f"Template não encontrado: {SRC}")

    prs = Presentation(str(SRC))
    delete_all_slides(prs)
    W, H = prs.slide_width, prs.slide_height

    # 1 — Capa
    s = new_slide(prs)
    add_num_badge(s, "04")
    _, tf = add_textbox(s, Inches(0.7), Inches(1.35), Inches(8.6), Inches(1.9))
    write_lines(
        tf,
        [
            ("Aula 04: A Linha do Tempo das", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 6}),
            ("Plataformas e as Restrições Técnicas", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 16}),
            ("Histórico de hardware · Viewport retrô e stretch clássico", {"size": 14, "color": TEXT_DIM, "font": BODY_FONT}),
        ],
    )
    add_panel(s, Inches(0.7), Inches(3.7), Inches(4.4), Inches(0.9))
    _, tf = add_textbox(s, Inches(0.9), Inches(3.85), Inches(4.0), Inches(0.7))
    write_lines(
        tf,
        [
            ("MÓDULO 1", {"size": 12, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 2}),
            ("FUNDAÇÕES, CULTURA E INTERFACE", {"size": 11, "color": TEXT, "font": BODY_FONT}),
        ],
    )
    add_footer(s, W, H)

    # 2 — Onde estamos
    s = new_slide(prs)
    add_num_badge(s, "01")
    add_title(s, "Onde estamos na trilha")
    steps = [
        ("AULA 02", "Player anda\n(Input Map + script)"),
        ("AULA 03", "Nearest + máscara\ncultural"),
        ("AULA 04", "Quadro / resolução\n/ plataforma"),
    ]
    for i, (k, v) in enumerate(steps):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.4), Inches(3.0), Inches(2.5))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.65), Inches(2.6), Inches(2.1))
        lines = [
            (k, {"size": 14, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
        ]
        for part in v.split("\n"):
            lines.append((part, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 2}))
        write_lines(tf, lines)
        if i < 2:
            _, tf = add_textbox(s, left + Inches(2.7), Inches(2.4), Inches(0.6), Inches(0.4))
            write_lines(tf, [("→", {"size": 20, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER})])
    add_footer(s, W, H)

    # 3 — Ementa
    s = new_slide(prs)
    add_num_badge(s, "02")
    add_title(s, "Tópico da ementa")
    add_panel(s, Inches(0.55), Inches(1.4), Inches(8.9), Inches(2.8))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.9), Inches(8.2), Inches(2.0))
    write_lines(
        tf,
        [
            ("Histórico dos jogos e suas", {"size": 22, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("plataformas de hardware", {"size": 22, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Da restrição técnica à criatividade — e ao contrato de tela na Godot.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 4 — O que é plataforma
    s = new_slide(prs)
    add_num_badge(s, "03")
    add_title(s, "O que é plataforma")
    cards = [
        ("HARDWARE", "CPU, memória, vídeo, entrada — o corpo da máquina."),
        ("CONTRATO", "O que o software pode pedir: pixels, cores, sprites."),
        ("GERAÇÃO", "Cada era redefine o orçamento visual do designer."),
    ]
    for i, (k, v) in enumerate(cards):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.3), Inches(3.0), Inches(2.8))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.6), Inches(2.6), Inches(2.3))
        write_lines(
            tf,
            [
                (k, {"size": 14, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
                (v, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
    add_footer(s, W, H)

    # 5 — Linha do tempo (mesa)
    s = new_slide(prs)
    add_num_badge(s, "04")
    add_title(s, "Linha do tempo — mesa")
    eras = [
        ("ATARI 2600", "~160×192 · poucos objetos por linha."),
        ("NES", "256×240 · 8 sprites/scanline · paleta curta."),
        ("SNES", "256×224 · mais cores/camadas · Mode 7."),
        ("HD / HOJE", "720p→4K · o limite muda de forma."),
    ]
    for i, (k, v) in enumerate(eras):
        col = i % 2
        row = i // 2
        left = Inches(0.55 + col * 4.6)
        top = Inches(1.15 + row * 1.7)
        add_panel(s, left, top, Inches(4.35), Inches(1.5))
        _, tf = add_textbox(s, left + Inches(0.25), top + Inches(0.25), Inches(3.9), Inches(1.1))
        write_lines(
            tf,
            [
                (k, {"size": 14, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 8}),
                (v, {"size": 13, "color": TEXT}),
            ],
        )
    add_footer(s, W, H)

    # 6 — Linha do tempo (portátil)
    s = new_slide(prs)
    add_num_badge(s, "05")
    add_title(s, "Linha do tempo — portátil")
    portables = [
        ("GAME BOY", "160×144 · 4 tons · silhueta > detalhe."),
        ("GBA", "240×160 · mais cor, tela perto do rosto."),
        ("DS / MODERNOS", "Telas dual / HD · densidade de UI importa."),
    ]
    for i, (k, v) in enumerate(portables):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.4), Inches(3.0), Inches(2.6))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.7), Inches(2.6), Inches(2.1))
        write_lines(
            tf,
            [
                (k, {"size": 14, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
                (v, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
    add_footer(s, W, H)

    # 7 — Restrições
    s = new_slide(prs)
    add_num_badge(s, "06")
    add_title(s, "Restrições que moldam o design")
    traits = [
        ("RESOLUÇÃO", "Poucos pixels — cada um conta."),
        ("PALETA", "Poucas cores por tile/sprite."),
        ("SPRITES / SCANLINE", "Limite por linha → flicker e truques."),
        ("MEMÓRIA", "Tiles reutilizados viram linguagem."),
    ]
    for i, (k, v) in enumerate(traits):
        top = Inches(1.1 + i * 0.8)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.68))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.12), Inches(8.4), Inches(0.5))
        write_lines(tf, [(f"{k}  —  {v}", {"size": 15, "color": TEXT})])
    add_footer(s, W, H)

    # 8 — Criatividade sob limite
    s = new_slide(prs)
    add_num_badge(s, "07")
    add_title(s, "Criatividade sob limite")
    examples = [
        ("FLICKER (NES)", "Sprites reordenados quando passam do limite por scanline."),
        ("TILES", "O mesmo bloco monta biomas inteiros."),
        ("MODE 7 (SNES)", "Perspectiva falsa e ondas — sem 3D de verdade."),
    ]
    for i, (k, v) in enumerate(examples):
        top = Inches(1.2 + i * 1.15)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(1.0))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.18), Inches(8.4), Inches(0.7))
        write_lines(
            tf,
            [
                (k, {"size": 13, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 4}),
                (v, {"size": 14, "color": TEXT}),
            ],
        )
    add_footer(s, W, H)

    # 9 — Retrô hoje
    s = new_slide(prs)
    add_num_badge(s, "08")
    add_title(s, "Retrô hoje — contratos 16:9")
    add_panel(s, Inches(0.55), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(0.8), Inches(1.55), Inches(3.8), Inches(2.6))
    write_lines(
        tf,
        [
            ("320 × 180", {"size": 22, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Pixels grandes e legíveis.", {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("Contrato “8/16-bit moderno”.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_panel(s, Inches(5.15), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(5.4), Inches(1.55), Inches(3.8), Inches(2.6))
    write_lines(
        tf,
        [
            ("480 × 270", {"size": 22, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Mesmo aspect, mais detalhe.", {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("Ainda claramente retrô.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 10 — Ponte prática
    s = new_slide(prs)
    add_num_badge(s, "09")
    add_title(s, "Ponte: do conceito à oficina")
    _, tf = add_textbox(s, Inches(0.7), Inches(1.4), Inches(8.6), Inches(2.8))
    write_lines(
        tf,
        [
            ("Hoje configuramos o contrato de tela.", {"size": 22, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Viewport baixo + stretch viewport/keep = pixels limpos.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Nearest (Aula 03) protege o asset; stretch protege o quadro.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 16}),
            ("Próximo: Project Settings → Display → Window.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 11 — Viewport vs janela
    s = new_slide(prs)
    add_num_badge(s, "10")
    add_title(s, "Viewport vs janela")
    add_panel(s, Inches(0.55), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(0.8), Inches(1.5), Inches(3.8), Inches(2.6))
    write_lines(
        tf,
        [
            ("VIEWPORT", {"size": 16, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Resolução de design.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Viewport Width / Height — o quadro nativo do jogo.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_panel(s, Inches(5.15), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(5.4), Inches(1.5), Inches(3.8), Inches(2.6))
    write_lines(
        tf,
        [
            ("JANELA", {"size": 16, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Tamanho na tela.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Window Override — ex.: 1280×720 com viewport 320×180.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 12 — Stretch Mode viewport
    s = new_slide(prs)
    add_num_badge(s, "11")
    add_title(s, "Stretch Mode = viewport")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.6), Inches(8.2), Inches(2.6))
    write_lines(
        tf,
        [
            ("Renderiza no tamanho base — depois escala para a janela.", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("É o caminho clássico para pixel art: o jogo “pensa” em poucos pixels.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Project Settings → Display → Window → Stretch → Mode → viewport", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 13 — Aspect keep + integer
    s = new_slide(prs)
    add_num_badge(s, "12")
    add_title(s, "Aspect keep + Scale integer")
    add_panel(s, Inches(0.55), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(0.8), Inches(1.5), Inches(3.8), Inches(2.6))
    write_lines(
        tf,
        [
            ("ASPECT KEEP", {"size": 16, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Mantém a proporção.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Barras pretas se a janela não for 16:9.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_panel(s, Inches(5.15), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(5.4), Inches(1.5), Inches(3.8), Inches(2.6))
    write_lines(
        tf,
        [
            ("SCALE INTEGER", {"size": 16, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Pixels inteiros.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Evita escala fracionária que “quebra” a grade.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 14 — Checklist
    s = new_slide(prs)
    add_num_badge(s, "13")
    add_title(s, "Checklist do artefato")
    checks = [
        "Viewport Width×Height = 320×180 ou 480×270.",
        "Stretch Mode = viewport.",
        "Stretch Aspect = keep.",
        "Stretch Scale Mode = integer.",
        "Window Override definido (ex.: 1280×720).",
        "Play: pixels nítidos + proporção ao redimensionar.",
        "Nearest ativo (eco Aula 03) + anotações enviadas.",
    ]
    for i, text in enumerate(checks):
        top = Inches(1.0 + i * 0.48)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.42))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.05), Inches(8.4), Inches(0.32))
        write_lines(tf, [(f"□  {text}", {"size": 12, "color": TEXT})])
    add_footer(s, W, H)

    # 15 — Fechamento
    s = new_slide(prs)
    add_num_badge(s, "14")
    add_title(s, "Fechamento e Altar")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.55), Inches(8.2), Inches(2.7))
    write_lines(
        tf,
        [
            ("FIM DA AULA 04: GUARDIÃO DA RESOLUÇÃO", {"size": 15, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Toda plataforma escreve um contrato invisível de pixels e cores.", {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Envie as anotações. Leve o código ao Altar quando o Mestre liberar.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Próxima trilha: ClassInd, IARC e design saudável (sem Godot).", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    return prs


def export_pdf_with_powerpoint(pptx_path: Path, pdf_path: Path) -> bool:
    """Exporta PDF via Microsoft PowerPoint (Windows COM)."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        ps = f"""
$ppt = New-Object -ComObject PowerPoint.Application
$ppt.Visible = [Microsoft.Office.Core.MsoTriState]::msoTrue
$pres = $ppt.Presentations.Open('{pptx_path}', $true, $false, $false)
$pres.SaveAs('{pdf_path}', 32)
$pres.Close()
$ppt.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($pres) | Out-Null
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null
"""
        import subprocess

        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            print(result.stdout)
            print(result.stderr, file=sys.stderr)
            return False
        return pdf_path.exists()

    powerpoint = win32com.client.Dispatch("PowerPoint.Application")
    powerpoint.Visible = 1
    presentation = powerpoint.Presentations.Open(str(pptx_path), WithWindow=False)
    presentation.SaveAs(str(pdf_path), 32)
    presentation.Close()
    powerpoint.Quit()
    return pdf_path.exists()


def main():
    prs = build_deck()
    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"PPTX: {OUT_PPTX} ({OUT_PPTX.stat().st_size} bytes)")

    if OUT_PDF.exists():
        OUT_PDF.unlink()
    ok = export_pdf_with_powerpoint(OUT_PPTX.resolve(), OUT_PDF.resolve())
    if ok:
        print(f"PDF:  {OUT_PDF} ({OUT_PDF.stat().st_size} bytes)")
    else:
        print("AVISO: falha ao exportar PDF via PowerPoint.", file=sys.stderr)
        sys.exit(2)

    check = Presentation(str(OUT_PPTX))
    print(f"Slides: {len(check.slides)}")


if __name__ == "__main__":
    main()
