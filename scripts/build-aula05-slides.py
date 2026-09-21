#!/usr/bin/env python3
"""Gera os slides da Aula 05 a partir do template visual da Aula 01.

Saídas:
  assets/docs/aulas/aula05_classind_iarc_slides.pptx
  assets/docs/aulas/aula05_classind_iarc_slides.pdf  (via PowerPoint COM, se disponível)
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pptx.oxml import parse_xml

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "assets" / "docs" / "aulas" / "aula01_godot_slides.pptx"
OUT_PPTX = ROOT / "assets" / "docs" / "aulas" / "aula05_classind_iarc_slides.pptx"
OUT_PDF = ROOT / "assets" / "docs" / "aulas" / "aula05_classind_iarc_slides.pdf"
COVERS = ROOT / "assets" / "classind-dle" / "covers"
FAIXAS_DIAGRAM = ROOT / "assets" / "slides" / "diagrama-das-faixas-etarias.png"

BG = RGBColor(0x0D, 0x0A, 0x10)
PANEL = RGBColor(0x18, 0x12, 0x1A)
GOLD = RGBColor(0xCF, 0xA7, 0x59)
GOLD_BRIGHT = RGBColor(0xF2, 0xD5, 0x9A)
GOLD_DARK = RGBColor(0x7A, 0x5C, 0x2E)
TEXT = RGBColor(0xEC, 0xE1, 0xD1)
TEXT_DIM = RGBColor(0xAB, 0x9C, 0x8A)
BLOOD = RGBColor(0x7A, 0x1F, 0x2B)

TITLE_FONT = "Cinzel"
BODY_FONT = "Georgia"
FOOTER = "AULA 05 | MÓDULO 1"


def _set_run(run, *, size_pt, bold=False, color=TEXT, font=BODY_FONT):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "cs", "ea"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = parse_xml(
                f'<a:{tag} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'
            )
            rPr.append(el)
        else:
            el.set("typeface", font)


def _clear_paragraphs(tf):
    p = tf.paragraphs[0]
    p.clear()
    for extra in list(tf.paragraphs)[1:]:
        p_elem = extra._p
        p_elem.getparent().remove(p_elem)
    return tf.paragraphs[0]


def add_textbox(slide, left, top, width, height, *, word_wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    return box, tf


def write_lines(tf, lines, *, size_pt=18, bold=False, color=TEXT, font=BODY_FONT, align=PP_ALIGN.LEFT, space_after=6):
    first = True
    for item in lines:
        if isinstance(item, tuple):
            text, opts = item[0], item[1]
        else:
            text, opts = item, {}
        p = _clear_paragraphs(tf) if first else tf.add_paragraph()
        first = False
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", space_after))
        run = p.add_run()
        run.text = text
        _set_run(
            run,
            size_pt=opts.get("size", size_pt),
            bold=opts.get("bold", bold),
            color=opts.get("color", color),
            font=opts.get("font", font),
        )


def paint_background(slide, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD_DARK
    bar.line.fill.background()
    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), height - Inches(0.08), width, Inches(0.08)
    )
    foot.fill.solid()
    foot.fill.fore_color.rgb = GOLD_DARK
    foot.line.fill.background()


def add_footer(slide, width, height, text=FOOTER):
    _, tf = add_textbox(
        slide,
        Inches(0.5),
        height - Inches(0.45),
        width - Inches(1.0),
        Inches(0.3),
    )
    write_lines(tf, [text], size_pt=10, color=TEXT_DIM, font=TITLE_FONT, align=PP_ALIGN.LEFT)


def add_num_badge(slide, number: str):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(0.28), Inches(0.7), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLOOD
    badge.line.color.rgb = GOLD
    badge.line.width = Pt(1)
    tf = badge.text_frame
    tf.word_wrap = False
    p = _clear_paragraphs(tf)
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = number
    _set_run(run, size_pt=14, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_title(slide, title: str, left=Inches(1.35), top=Inches(0.28), width=Inches(8.2)):
    _, tf = add_textbox(slide, left, top, width, Inches(0.55))
    write_lines(tf, [title], size_pt=22, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_panel(slide, left, top, width, height):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = GOLD_DARK
    panel.line.width = Pt(1.25)
    return panel


def delete_all_slides(prs: Presentation):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def new_slide(prs: Presentation):
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    for shape in list(slide.shapes):
        sp = shape._element
        sp.getparent().remove(sp)
    paint_background(slide, prs.slide_width, prs.slide_height)
    return slide


def add_cover(slide, filename: str, left, top, width, height):
    """Insere capa WebP convertendo para PNG temporário (python-pptx não aceita WebP)."""
    path = COVERS / filename
    if not path.exists():
        return None
    try:
        from PIL import Image
    except ImportError:
        print(f"AVISO: Pillow ausente — pulando capa {filename}", file=sys.stderr)
        return None

    tmp = ROOT / "assets" / "docs" / "aulas" / "_tmp_slide_covers"
    tmp.mkdir(parents=True, exist_ok=True)
    png_path = tmp / (path.stem + ".png")
    if not png_path.exists() or png_path.stat().st_mtime < path.stat().st_mtime:
        with Image.open(path) as img:
            img.convert("RGB").save(png_path, "PNG")
    return slide.shapes.add_picture(str(png_path), left, top, width=width, height=height)


def add_slide_image(slide, path: Path, left, top, width, height):
    if not path.exists():
        print(f"AVISO: imagem ausente — {path}", file=sys.stderr)
        return None
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def build_deck() -> Presentation:
    if not SRC.exists():
        raise SystemExit(f"Template não encontrado: {SRC}")

    prs = Presentation(str(SRC))
    delete_all_slides(prs)
    W, H = prs.slide_width, prs.slide_height

    # 1 — Capa
    s = new_slide(prs)
    add_num_badge(s, "05")
    _, tf = add_textbox(s, Inches(0.7), Inches(1.35), Inches(8.6), Inches(1.9))
    write_lines(
        tf,
        [
            ("Aula 05: Classificação Indicativa", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 6}),
            ("(ClassInd), IARC e Design Saudável", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 16}),
            ("Faixas etárias · ClassInd-dle · Adequação de público", {"size": 14, "color": TEXT_DIM, "font": BODY_FONT}),
        ],
    )
    add_panel(s, Inches(0.7), Inches(3.7), Inches(4.4), Inches(0.9))
    _, tf = add_textbox(s, Inches(0.9), Inches(3.85), Inches(4.0), Inches(0.7))
    write_lines(
        tf,
        [
            ("MÓDULO 1", {"size": 12, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 2}),
            ("FUNDAÇÕES, CULTURA E INTERFACE", {"size": 11, "color": TEXT, "font": BODY_FONT}),
        ],
    )
    add_footer(s, W, H)

    # 2 — Onde estamos
    s = new_slide(prs)
    add_num_badge(s, "01")
    add_title(s, "Onde estamos na trilha")
    steps = [
        ("AULA 04", "Contrato de tela\n(viewport / stretch)"),
        ("AULA 05", "Para quem publica?\nClassInd + IARC"),
        ("PRÁTICA", "ClassInd-dle +\nPatch Note no site"),
    ]
    for i, (k, v) in enumerate(steps):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.4), Inches(3.0), Inches(2.5))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.65), Inches(2.6), Inches(2.1))
        lines = [
            (k, {"size": 14, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
        ]
        for part in v.split("\n"):
            lines.append((part, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 2}))
        write_lines(tf, lines)
        if i < 2:
            _, tf = add_textbox(s, left + Inches(2.7), Inches(2.4), Inches(0.6), Inches(0.4))
            write_lines(tf, [("→", {"size": 20, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER})])
    add_footer(s, W, H)

    # 3 — Ementa
    s = new_slide(prs)
    add_num_badge(s, "02")
    add_title(s, "Tópico da ementa")
    add_panel(s, Inches(0.55), Inches(1.4), Inches(8.9), Inches(2.8))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.9), Inches(8.2), Inches(2.0))
    write_lines(
        tf,
        [
            ("Sistemas de classificação indicativa", {"size": 22, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("e adequação de público", {"size": 22, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("O selo não é só marketing — é consequência de feedbacks de design.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 4 — ClassInd / IARC / Design
    s = new_slide(prs)
    add_num_badge(s, "03")
    add_title(s, "Três ideias da aula")
    cards = [
        ("CLASSIND", "Aviso brasileiro de conteúdo — consumidor, lojas e responsabilidade do dev."),
        ("IARC", "Formulário único → selos multi-região/loja, gratuito no digital."),
        ("DESIGN SAUDÁVEL", "Preservar o loop-core sem inflar violência, sexo ou drogas."),
    ]
    for i, (k, v) in enumerate(cards):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.3), Inches(3.0), Inches(2.8))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.6), Inches(2.6), Inches(2.3))
        write_lines(
            tf,
            [
                (k, {"size": 13, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
                (v, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
    add_footer(s, W, H)

    # 5 — Faixas (diagrama)
    s = new_slide(prs)
    add_num_badge(s, "04")
    add_title(s, "Faixas — leitura operacional")
    if FAIXAS_DIAGRAM.exists():
        # Diagrama em paisagem, centralizado sob o título
        add_slide_image(s, FAIXAS_DIAGRAM, Inches(0.7), Inches(1.15), Inches(8.6), Inches(3.55))
    else:
        faixa_rows = [
            ("L", "Livre — fantasia/comicidade; sem gore realista."),
            ("10", "Violência leve; medo leve; sem carga sexual."),
            ("12–14", "Mais presença / intensidade; temas e linguagem sobem."),
            ("16–18", "Violência forte/extrema; sexo; drogas com destaque."),
        ]
        for i, (k, v) in enumerate(faixa_rows):
            top = Inches(1.1 + i * 0.8)
            add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.68))
            _, tf = add_textbox(s, Inches(0.8), top + Inches(0.12), Inches(8.4), Inches(0.5))
            write_lines(tf, [(f"{k}  —  {v}", {"size": 15, "color": TEXT})])
    add_footer(s, W, H)

    # 6 — Três eixos
    s = new_slide(prs)
    add_num_badge(s, "05")
    add_title(s, "Três eixos clássicos")
    axes = [
        ("VIOLÊNCIA", "Contra quem? Sangue? Cadáveres? Fantasia vs realismo."),
        ("SEXO", "Insinuação, nudez, atos, recompensa sexual."),
        ("DROGAS", "Menção, uso como mecânica, glamourização."),
    ]
    for i, (k, v) in enumerate(axes):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.4), Inches(3.0), Inches(2.6))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.7), Inches(2.6), Inches(2.1))
        write_lines(
            tf,
            [
                (k, {"size": 14, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
                (v, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
    add_footer(s, W, H)

    # 7 — Atenuantes / agravantes
    s = new_slide(prs)
    add_num_badge(s, "06")
    add_title(s, "Atenuantes × agravantes")
    add_panel(s, Inches(0.55), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(0.8), Inches(1.5), Inches(3.8), Inches(2.6))
    write_lines(
        tf,
        [
            ("ATENUAM", {"size": 16, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Fantasia · não-humano · comicidade", {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("Sem sangue · cura simbólica / mágica", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_panel(s, Inches(5.15), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(5.4), Inches(1.5), Inches(3.8), Inches(2.6))
    write_lines(
        tf,
        [
            ("AGRAVAM", {"size": 16, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Realismo · gore · cadáveres", {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("Nudez · glamourização de substâncias", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 8 — IARC
    s = new_slide(prs)
    add_num_badge(s, "07")
    add_title(s, "IARC — selos a partir de um formulário")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.6), Inches(8.2), Inches(2.6))
    write_lines(
        tf,
        [
            ("Você declara o conteúdo → o consórcio devolve selos.", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Gratuito e pensado para lojas digitais. As perguntas espelham os eixos.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Mentir no formulário não é truque — é risco comercial e ético.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 9 — Ponte prática
    s = new_slide(prs)
    add_num_badge(s, "08")
    add_title(s, "Ponte: do conceito à oficina")
    _, tf = add_textbox(s, Inches(0.7), Inches(1.4), Inches(8.6), Inches(2.8))
    write_lines(
        tf,
        [
            ("Mesmo loop, faixas diferentes — o feedback decide.", {"size": 20, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Parte 1: ClassInd-dle (Higher/Lower ao vivo).", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Parte 2: higienizar um pitch 16+/18+ até Livre ou 10.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 16}),
            ("Sem Godot hoje — a oficina é o site.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 10 — ClassInd-dle (passos)
    s = new_slide(prs)
    add_num_badge(s, "09")
    add_title(s, "Parte 1 — ClassInd-dle")
    steps_dle = [
        "Mestre cria a sala e mostra o código no telão.",
        "Alunos entram pelo CTA da Oficina (Mestre não vota).",
        "Compare A vs B: qual exige a idade mais alta?",
        "Vote uma vez. Placar ao vivo sem spoiler.",
        "Após o revelar: faixas + acerto/erro + placar pessoal.",
    ]
    for i, text in enumerate(steps_dle):
        top = Inches(1.1 + i * 0.58)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.5))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.08), Inches(8.4), Inches(0.35))
        write_lines(tf, [(f"{i + 1}.  {text}", {"size": 14, "color": TEXT})])
    add_footer(s, W, H)

    # 10b — Exemplo visual Higher/Lower (capas)
    s = new_slide(prs)
    add_num_badge(s, "09b")
    add_title(s, "Exemplo de comparação (capas)")
    add_cover(s, "the-sims-4.webp", Inches(1.0), Inches(1.2), Inches(2.6), Inches(3.45))
    add_cover(s, "hollow-knight.webp", Inches(6.4), Inches(1.2), Inches(2.6), Inches(3.45))
    _, tf = add_textbox(s, Inches(3.7), Inches(2.4), Inches(2.4), Inches(1.0))
    write_lines(
        tf,
        [
            ("VS", {"size": 28, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("Pegadinha: Sims 12 > HK 10", {"size": 12, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 11 — Adequação reversa
    s = new_slide(prs)
    add_num_badge(s, "10")
    add_title(s, "Parte 2 — Adequação reversa")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(5.5), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.8), Inches(1.5), Inches(5.0), Inches(2.7))
    write_lines(
        tf,
        [
            ("Missão: pitch 16+/18+ → Livre ou 10.", {"size": 16, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 10}),
            ("Gere um desafio procedural no wizard.", {"size": 14, "color": TEXT, "space_after": 8}),
            ("Reescreva visual / narrativa / cura / inimigos.", {"size": 14, "color": TEXT, "space_after": 8}),
            ("Preserve o loop-core. Finalize com Patch Note.", {"size": 14, "color": TEXT, "space_after": 8}),
        ],
    )
    add_cover(s, "doom-eternal.webp", Inches(6.4), Inches(1.35), Inches(2.7), Inches(3.0))
    add_footer(s, W, H)

    # 12 — Pitches
    s = new_slide(prs)
    add_num_badge(s, "11")
    add_title(s, "Pitches da mesa")
    pitches = [
        ("NECRÓPOLE VIRAL", "Gore / infecção → robôs, faíscas, baterias."),
        ("SOMBRA DO CONTRATO", "Ritual / sangue → silhuetas e pacto simbólico."),
        ("APP DE DESTINOS", "Nudez / sexo → afeto sem nudez explícita."),
        ("PORÃO DAS HORAS", "Drogas realistas → amuleto / luz ritual."),
    ]
    for i, (k, v) in enumerate(pitches):
        col = i % 2
        row = i // 2
        left = Inches(0.55 + col * 4.6)
        top = Inches(1.15 + row * 1.7)
        add_panel(s, left, top, Inches(4.35), Inches(1.5))
        _, tf = add_textbox(s, left + Inches(0.25), top + Inches(0.25), Inches(3.9), Inches(1.1))
        write_lines(
            tf,
            [
                (k, {"size": 13, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 8}),
                (v, {"size": 13, "color": TEXT}),
            ],
        )
    add_footer(s, W, H)

    # 13 — Checklist
    s = new_slide(prs)
    add_num_badge(s, "12")
    add_title(s, "Checklist do artefato")
    checks = [
        "Participei do ClassInd-dle com pelo menos um voto.",
        "Anotei 1 insight (detalhe que mudou a faixa).",
        "Patch Note: original → higienizado → argumentos → L ou 10.",
        "Mecânica-core preservada (o verbo do jogo).",
        "Anotações + síntese enviadas na página da aula.",
        "Código no Altar quando o Mestre liberar.",
    ]
    for i, text in enumerate(checks):
        top = Inches(1.1 + i * 0.52)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.45))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.06), Inches(8.4), Inches(0.35))
        write_lines(tf, [(f"□  {text}", {"size": 13, "color": TEXT})])
    add_footer(s, W, H)

    # 14 — Fechamento
    s = new_slide(prs)
    add_num_badge(s, "13")
    add_title(s, "Fechamento e Altar")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.55), Inches(8.2), Inches(2.7))
    write_lines(
        tf,
        [
            ("FIM DA AULA 05: GUARDIÃO DA FAIXA", {"size": 15, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Higienizar é redesenhar o que o jogador vê e sente — não apagar o jogo.", {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Envie o Patch Note pelo wizard. Código ao Altar quando o Mestre liberar.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Capas: assets/classind-dle/covers/INVENTARIO.md · diagrama de faixas já nos slides.", {"size": 12, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    return prs


def export_pdf_with_powerpoint(pptx_path: Path, pdf_path: Path) -> bool:
    """Exporta PDF via Microsoft PowerPoint (Windows COM)."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        ps = f"""
$ppt = New-Object -ComObject PowerPoint.Application
$ppt.Visible = [Microsoft.Office.Core.MsoTriState]::msoTrue
$pres = $ppt.Presentations.Open('{pptx_path}', $true, $false, $false)
$pres.SaveAs('{pdf_path}', 32)
$pres.Close()
$ppt.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($pres) | Out-Null
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null
"""
        import subprocess

        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            print(result.stdout)
            print(result.stderr, file=sys.stderr)
            return False
        return pdf_path.exists()

    powerpoint = win32com.client.Dispatch("PowerPoint.Application")
    powerpoint.Visible = 1
    presentation = powerpoint.Presentations.Open(str(pptx_path), WithWindow=False)
    presentation.SaveAs(str(pdf_path), 32)
    presentation.Close()
    powerpoint.Quit()
    return pdf_path.exists()


def main():
    prs = build_deck()
    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"PPTX: {OUT_PPTX} ({OUT_PPTX.stat().st_size} bytes)")

    if OUT_PDF.exists():
        OUT_PDF.unlink()
    ok = export_pdf_with_powerpoint(OUT_PPTX.resolve(), OUT_PDF.resolve())
    if ok:
        print(f"PDF:  {OUT_PDF} ({OUT_PDF.stat().st_size} bytes)")
    else:
        print("AVISO: falha ao exportar PDF via PowerPoint.", file=sys.stderr)
        sys.exit(2)

    check = Presentation(str(OUT_PPTX))
    print(f"Slides: {len(check.slides)}")


if __name__ == "__main__":
    main()
