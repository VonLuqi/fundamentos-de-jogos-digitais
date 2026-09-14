#!/usr/bin/env python3
"""Gera os slides da Aula 03 a partir do template visual da Aula 01.

Saídas:
  assets/docs/aulas/aula03_homo_ludens_slides.pptx
  assets/docs/aulas/aula03_homo_ludens_slides.pdf  (via PowerPoint COM, se disponível)
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pptx.oxml import parse_xml

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "assets" / "docs" / "aulas" / "aula01_godot_slides.pptx"
OUT_PPTX = ROOT / "assets" / "docs" / "aulas" / "aula03_homo_ludens_slides.pptx"
OUT_PDF = ROOT / "assets" / "docs" / "aulas" / "aula03_homo_ludens_slides.pdf"

BG = RGBColor(0x0D, 0x0A, 0x10)
PANEL = RGBColor(0x18, 0x12, 0x1A)
GOLD = RGBColor(0xCF, 0xA7, 0x59)
GOLD_BRIGHT = RGBColor(0xF2, 0xD5, 0x9A)
GOLD_DARK = RGBColor(0x7A, 0x5C, 0x2E)
TEXT = RGBColor(0xEC, 0xE1, 0xD1)
TEXT_DIM = RGBColor(0xAB, 0x9C, 0x8A)
BLOOD = RGBColor(0x7A, 0x1F, 0x2B)

TITLE_FONT = "Cinzel"
BODY_FONT = "Georgia"
FOOTER = "AULA 03 | MÓDULO 1"


def _set_run(run, *, size_pt, bold=False, color=TEXT, font=BODY_FONT):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "cs", "ea"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = parse_xml(
                f'<a:{tag} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'
            )
            rPr.append(el)
        else:
            el.set("typeface", font)


def _clear_paragraphs(tf):
    p = tf.paragraphs[0]
    p.clear()
    for extra in list(tf.paragraphs)[1:]:
        p_elem = extra._p
        p_elem.getparent().remove(p_elem)
    return tf.paragraphs[0]


def add_textbox(slide, left, top, width, height, *, word_wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    return box, tf


def write_lines(tf, lines, *, size_pt=18, bold=False, color=TEXT, font=BODY_FONT, align=PP_ALIGN.LEFT, space_after=6):
    first = True
    for item in lines:
        if isinstance(item, tuple):
            text, opts = item[0], item[1]
        else:
            text, opts = item, {}
        p = _clear_paragraphs(tf) if first else tf.add_paragraph()
        first = False
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", space_after))
        run = p.add_run()
        run.text = text
        _set_run(
            run,
            size_pt=opts.get("size", size_pt),
            bold=opts.get("bold", bold),
            color=opts.get("color", color),
            font=opts.get("font", font),
        )


def paint_background(slide, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD_DARK
    bar.line.fill.background()
    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), height - Inches(0.08), width, Inches(0.08)
    )
    foot.fill.solid()
    foot.fill.fore_color.rgb = GOLD_DARK
    foot.line.fill.background()


def add_footer(slide, width, height, text=FOOTER):
    _, tf = add_textbox(
        slide,
        Inches(0.5),
        height - Inches(0.45),
        width - Inches(1.0),
        Inches(0.3),
    )
    write_lines(tf, [text], size_pt=10, color=TEXT_DIM, font=TITLE_FONT, align=PP_ALIGN.LEFT)


def add_num_badge(slide, number: str):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(0.28), Inches(0.7), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLOOD
    badge.line.color.rgb = GOLD
    badge.line.width = Pt(1)
    tf = badge.text_frame
    tf.word_wrap = False
    p = _clear_paragraphs(tf)
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = number
    _set_run(run, size_pt=14, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_title(slide, title: str, left=Inches(1.35), top=Inches(0.28), width=Inches(8.2)):
    _, tf = add_textbox(slide, left, top, width, Inches(0.55))
    write_lines(tf, [title], size_pt=22, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_panel(slide, left, top, width, height):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = GOLD_DARK
    panel.line.width = Pt(1.25)
    return panel


def delete_all_slides(prs: Presentation):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def new_slide(prs: Presentation):
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    for shape in list(slide.shapes):
        sp = shape._element
        sp.getparent().remove(sp)
    paint_background(slide, prs.slide_width, prs.slide_height)
    return slide


def build_deck() -> Presentation:
    if not SRC.exists():
        raise SystemExit(f"Template não encontrado: {SRC}")

    prs = Presentation(str(SRC))
    delete_all_slides(prs)
    W, H = prs.slide_width, prs.slide_height

    # 1 — Capa
    s = new_slide(prs)
    add_num_badge(s, "03")
    _, tf = add_textbox(s, Inches(0.7), Inches(1.5), Inches(8.6), Inches(1.6))
    write_lines(
        tf,
        [
            ("Aula 03: Homo Ludens, Identidade", {"size": 26, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 8}),
            ("e Expressão Cultural", {"size": 26, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 16}),
            ("Cultura como jogo · Pixel art, importação e herói brasileiro", {"size": 14, "color": TEXT_DIM, "font": BODY_FONT}),
        ],
    )
    add_panel(s, Inches(0.7), Inches(3.7), Inches(4.4), Inches(0.9))
    _, tf = add_textbox(s, Inches(0.9), Inches(3.85), Inches(4.0), Inches(0.7))
    write_lines(
        tf,
        [
            ("MÓDULO 1", {"size": 12, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 2}),
            ("FUNDAÇÕES, CULTURA E INTERFACE", {"size": 11, "color": TEXT, "font": BODY_FONT}),
        ],
    )
    add_footer(s, W, H)

    # 2 — Onde estamos
    s = new_slide(prs)
    add_num_badge(s, "01")
    add_title(s, "Onde estamos na trilha")
    steps = [
        ("AULA 01", "Círculo Mágico + Inspector"),
        ("AULA 02", "Player anda (Input Map + script)"),
        ("AULA 03", "Máscara cultural + pixel nítido"),
    ]
    for i, (k, v) in enumerate(steps):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.4), Inches(3.0), Inches(2.5))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.7), Inches(2.6), Inches(2.0))
        write_lines(
            tf,
            [
                (k, {"size": 14, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
                (v, {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
        if i < 2:
            _, tf = add_textbox(s, left + Inches(2.7), Inches(2.4), Inches(0.6), Inches(0.4))
            write_lines(tf, [("→", {"size": 20, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER})])
    add_footer(s, W, H)

    # 3 — Ementa
    s = new_slide(prs)
    add_num_badge(s, "02")
    add_title(s, "Tópico da ementa")
    add_panel(s, Inches(0.55), Inches(1.4), Inches(8.9), Inches(2.8))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.9), Inches(8.2), Inches(2.0))
    write_lines(
        tf,
        [
            ("O jogo como elemento da cultura", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 16}),
            ("(Homo Ludens)", {"size": 20, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Do conceito de Huizinga à máscara do herói na Godot.", {"size": 15, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 4 — Huizinga
    s = new_slide(prs)
    add_num_badge(s, "03")
    add_title(s, "Huizinga — tese-chave")
    add_panel(s, Inches(0.55), Inches(1.2), Inches(8.9), Inches(3.2))
    _, tf = add_textbox(s, Inches(0.85), Inches(1.55), Inches(8.3), Inches(2.7))
    write_lines(
        tf,
        [
            ("Homo Ludens (1938)", {"size": 16, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 12}),
            ("A cultura humana surge e se desenvolve como jogo.", {"size": 20, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 14}),
            ("No jogo e pelo jogo a civilização toma forma — rito, linguagem, arte e competição carregam caráter lúdico.", {"size": 15, "color": TEXT, "space_after": 10}),
            ("O humano não é só quem pensa ou fabrica: é também quem joga.", {"size": 14, "color": TEXT_DIM}),
        ],
    )
    add_footer(s, W, H)

    # 5 — O que é jogar
    s = new_slide(prs)
    add_num_badge(s, "04")
    add_title(s, "O que é jogar — definição operacional")
    traits = [
        ("LIVRE", "Voluntário, não imposto."),
        ("LIMITES", "Tempo e espaço próprios."),
        ("REGRAS", "Aceitas por quem entra."),
        ("TENSÃO", "Prazer + risco lúdico."),
        ("FORA DO COTIDIANO", "Consciência de outro mundo."),
    ]
    for i, (k, v) in enumerate(traits):
        top = Inches(1.05 + i * 0.68)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.58))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.1), Inches(8.4), Inches(0.4))
        write_lines(
            tf,
            [
                (f"{k}  —  {v}", {"size": 14, "color": TEXT}),
            ],
        )
    add_footer(s, W, H)

    # 6 — Cultura jogada
    s = new_slide(prs)
    add_num_badge(s, "05")
    add_title(s, "Cultura jogada")
    cards = [
        ("RITO", "Cerimônia e festa como forma lúdica."),
        ("LINGUAGEM", "Jogo de sentidos e nomes."),
        ("ARTE", "Criação com regras e tensão."),
        ("COMPETIÇÃO", "Agón — disputa que gera ordem."),
    ]
    for i, (k, v) in enumerate(cards):
        col = i % 2
        row = i // 2
        left = Inches(0.55 + col * 4.6)
        top = Inches(1.15 + row * 1.7)
        add_panel(s, left, top, Inches(4.35), Inches(1.5))
        _, tf = add_textbox(s, left + Inches(0.25), top + Inches(0.25), Inches(3.9), Inches(1.1))
        write_lines(
            tf,
            [
                (k, {"size": 14, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 8}),
                (v, {"size": 14, "color": TEXT}),
            ],
        )
    add_footer(s, W, H)

    # 7 — Identidade nos jogos
    s = new_slide(prs)
    add_num_badge(s, "06")
    add_title(s, "Identidade nos jogos")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.85), Inches(1.55), Inches(8.3), Inches(2.7))
    write_lines(
        tf,
        [
            ("Criadores traduzem regionalidade em sistemas jogáveis.", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 14}),
            ("Não é “decorar” o jogo com cultura — é transformar lendas, fauna e paisagens em elementos interativos.", {"size": 15, "color": TEXT, "space_after": 10}),
            ("O sprite, o cenário e o loop dizem de onde falamos.", {"size": 15, "color": TEXT, "space_after": 10}),
            ("Hoje: o Player da Aula 02 veste uma máscara brasileira.", {"size": 14, "color": TEXT_DIM}),
        ],
    )
    add_footer(s, W, H)

    # 8 — Exemplos BR
    s = new_slide(prs)
    add_num_badge(s, "07")
    add_title(s, "Exemplos BR — inspiração")
    examples = [
        ("FOLCLORE", "Saci, Curupira, Cuca, Boitatá em indies nacionais."),
        ("REGIONAL", "Sertão, Amazônia, urbano contemporâneo."),
        ("FAUNA / SÍMBOLOS", "Motivos locais como identidade visual do herói."),
    ]
    for i, (k, v) in enumerate(examples):
        top = Inches(1.2 + i * 1.15)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(1.0))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.18), Inches(8.4), Inches(0.7))
        write_lines(
            tf,
            [
                (k, {"size": 13, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 4}),
                (v, {"size": 15, "color": TEXT}),
            ],
        )
    add_footer(s, W, H)

    # 9 — Ponte prática
    s = new_slide(prs)
    add_num_badge(s, "08")
    add_title(s, "Ponte: do conceito à oficina")
    _, tf = add_textbox(s, Inches(0.7), Inches(1.5), Inches(8.6), Inches(2.6))
    write_lines(
        tf,
        [
            ("Hoje o sprite vira máscara.", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Importar pixel art e deixar nítida (Nearest) é o ofício.", {"size": 16, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Escolher o que o herói representa é o gesto de Homo Ludens.", {"size": 16, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 16}),
            ("Próximo: pacote Tiny Hero e FileSystem.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 10 — Pacote Tiny Hero
    s = new_slide(prs)
    add_num_badge(s, "09")
    add_title(s, "Pacote Tiny Hero (CraftPix)")
    heroes = [
        ("PINK", "Pink_Monster.png"),
        ("OWLET", "Owlet_Monster.png"),
        ("DUDE", "Dude_Monster.png"),
    ]
    for i, (k, v) in enumerate(heroes):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.2), Inches(3.0), Inches(2.0))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.5), Inches(2.6), Inches(1.5))
        write_lines(
            tf,
            [
                (k, {"size": 16, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 10}),
                (v, {"size": 12, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
    _, tf = add_textbox(s, Inches(0.6), Inches(3.5), Inches(8.8), Inches(0.9))
    write_lines(
        tf,
        [
            ("MVP: use o PNG base (1 frame no Sprite2D).", {"size": 14, "color": TEXT, "space_after": 4}),
            ("Sheets Idle/Walk = exploração; animação completa fica para depois.", {"size": 13, "color": TEXT_DIM}),
        ],
    )
    add_footer(s, W, H)

    # 11 — Importar
    s = new_slide(prs)
    add_num_badge(s, "10")
    add_title(s, "Importar no FileSystem")
    steps = [
        "Crie a pasta sprites/hero no projeto da Aula 02.",
        "Arraste e solte o PNG escolhido no FileSystem.",
        "Confirme res://… e os metadados de importação.",
        "Atribua a textura ao Sprite2D do Player.",
    ]
    for i, text in enumerate(steps, start=1):
        top = Inches(1.1 + (i - 1) * 0.85)
        badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.55), top + Inches(0.1), Inches(0.5), Inches(0.5))
        badge.fill.solid()
        badge.fill.fore_color.rgb = BLOOD
        badge.line.color.rgb = GOLD
        tf = badge.text_frame
        p = _clear_paragraphs(tf)
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = f"{i:02d}"
        _set_run(run, size_pt=12, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)
        add_panel(s, Inches(1.25), top, Inches(8.1), Inches(0.7))
        _, tf = add_textbox(s, Inches(1.45), top + Inches(0.15), Inches(7.7), Inches(0.45))
        write_lines(tf, [text], size_pt=14, color=TEXT)
    add_footer(s, W, H)

    # 12 — Nearest
    s = new_slide(prs)
    add_num_badge(s, "11")
    add_title(s, "Nearest — pixel nítido")
    add_panel(s, Inches(0.55), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(0.8), Inches(1.5), Inches(3.8), Inches(2.6))
    write_lines(
        tf,
        [
            ("LINEAR", {"size": 18, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Suaviza e borra o pixel art.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Default em projetos novos — inimigo do pixel duro.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_panel(s, Inches(5.15), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(5.4), Inches(1.5), Inches(3.8), Inches(2.6))
    write_lines(
        tf,
        [
            ("NEAREST", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Preserva arestas de pixel.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Sprite2D → Texture → Filter, ou default do projeto.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 13 — Personalização cultural
    s = new_slide(prs)
    add_num_badge(s, "12")
    add_title(s, "Personalização cultural")
    anchors = [
        ("FOLCLORE", "Saci, Curupira, Iara, Boitatá, Cuca…"),
        ("FAUNA BR", "Onça, tucano, jabuti, capivara…"),
        ("URBANO / REGIONAL", "Cidade, cordel, frevo, sertão…"),
    ]
    for i, (k, v) in enumerate(anchors):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.2), Inches(3.0), Inches(2.2))
        _, tf = add_textbox(s, left + Inches(0.15), Inches(1.45), Inches(2.7), Inches(1.8))
        write_lines(
            tf,
            [
                (k, {"size": 13, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 10}),
                (v, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
    _, tf = add_textbox(s, Inches(0.6), Inches(3.6), Inches(8.8), Inches(0.8))
    write_lines(
        tf,
        [
            ("Nível 1: recolor ou crop + justificativa.  Nível 2: acessório/silhueta.  Nível 3: sheets (bônus).", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 14 — Checklist
    s = new_slide(prs)
    add_num_badge(s, "13")
    add_title(s, "Checklist do artefato")
    checks = [
        "ZIP Tiny Hero baixado e PNG no projeto.",
        "Sprite2D do Player com a textura importada.",
        "Filtro Nearest — pixel nítido no Play.",
        "Personalização cultural Nível 1+ (visual e/ou justificativa).",
        "CollisionShape ajustado + movimento da Aula 02 intacto.",
        "Anotações + síntese enviadas na página da aula.",
    ]
    for i, text in enumerate(checks):
        top = Inches(1.05 + i * 0.55)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.48))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.08), Inches(8.4), Inches(0.35))
        write_lines(tf, [(f"□  {text}", {"size": 13, "color": TEXT})])
    add_footer(s, W, H)

    # 15 — Fechamento
    s = new_slide(prs)
    add_num_badge(s, "14")
    add_title(s, "Fechamento e Altar")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.6), Inches(8.2), Inches(2.6))
    write_lines(
        tf,
        [
            ("FIM DA AULA 03: MÁSCARA DO HOMO LUDENS", {"size": 16, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("A cultura joga — e o seu herói agora carrega uma máscara brasileira.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Envie as anotações na página. Leve o código ao Altar quando o Mestre liberar.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Próxima trilha: câmera, cena de teste e primeira animação.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    return prs


def export_pdf_with_powerpoint(pptx_path: Path, pdf_path: Path) -> bool:
    """Exporta PDF via Microsoft PowerPoint (Windows COM)."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        ps = f"""
$ppt = New-Object -ComObject PowerPoint.Application
$ppt.Visible = [Microsoft.Office.Core.MsoTriState]::msoTrue
$pres = $ppt.Presentations.Open('{pptx_path}', $true, $false, $false)
$pres.SaveAs('{pdf_path}', 32)
$pres.Close()
$ppt.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($pres) | Out-Null
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null
"""
        import subprocess

        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            print(result.stdout)
            print(result.stderr, file=sys.stderr)
            return False
        return pdf_path.exists()

    powerpoint = win32com.client.Dispatch("PowerPoint.Application")
    powerpoint.Visible = 1
    presentation = powerpoint.Presentations.Open(str(pptx_path), WithWindow=False)
    presentation.SaveAs(str(pdf_path), 32)
    presentation.Close()
    powerpoint.Quit()
    return pdf_path.exists()


def main():
    prs = build_deck()
    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"PPTX: {OUT_PPTX} ({OUT_PPTX.stat().st_size} bytes)")

    if OUT_PDF.exists():
        OUT_PDF.unlink()
    ok = export_pdf_with_powerpoint(OUT_PPTX.resolve(), OUT_PDF.resolve())
    if ok:
        print(f"PDF:  {OUT_PDF} ({OUT_PDF.stat().st_size} bytes)")
    else:
        print("AVISO: falha ao exportar PDF via PowerPoint.", file=sys.stderr)
        sys.exit(2)

    check = Presentation(str(OUT_PPTX))
    print(f"Slides: {len(check.slides)}")


if __name__ == "__main__":
    main()
