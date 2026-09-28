#!/usr/bin/env python3
"""Gera os slides da Aula 06 a partir do template visual da Aula 01.

Saídas:
  assets/docs/aulas/aula06_mercado_loja_etica_slides.pptx
  assets/docs/aulas/aula06_mercado_loja_etica_slides.pdf  (via PowerPoint COM, se disponível)
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pptx.oxml import parse_xml

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "assets" / "docs" / "aulas" / "aula01_godot_slides.pptx"
OUT_PPTX = ROOT / "assets" / "docs" / "aulas" / "aula06_mercado_loja_etica_slides.pptx"
OUT_PDF = ROOT / "assets" / "docs" / "aulas" / "aula06_mercado_loja_etica_slides.pdf"

BG = RGBColor(0x0D, 0x0A, 0x10)
PANEL = RGBColor(0x18, 0x12, 0x1A)
GOLD = RGBColor(0xCF, 0xA7, 0x59)
GOLD_BRIGHT = RGBColor(0xF2, 0xD5, 0x9A)
GOLD_DARK = RGBColor(0x7A, 0x5C, 0x2E)
TEXT = RGBColor(0xEC, 0xE1, 0xD1)
TEXT_DIM = RGBColor(0xAB, 0x9C, 0x8A)
BLOOD = RGBColor(0x7A, 0x1F, 0x2B)

TITLE_FONT = "Cinzel"
BODY_FONT = "Georgia"
CODE_FONT = "Consolas"
FOOTER = "AULA 06 | MÓDULO 2"

HIERARQUIA_IMG = ROOT / "assets" / "docs" / "aulas" / "aula06-loja-etica" / "hierarquia-loja-godot.png"


def _set_run(run, *, size_pt, bold=False, color=TEXT, font=BODY_FONT):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "cs", "ea"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = parse_xml(
                f'<a:{tag} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'
            )
            rPr.append(el)
        else:
            el.set("typeface", font)


def _clear_paragraphs(tf):
    p = tf.paragraphs[0]
    p.clear()
    for extra in list(tf.paragraphs)[1:]:
        p_elem = extra._p
        p_elem.getparent().remove(p_elem)
    return tf.paragraphs[0]


def add_textbox(slide, left, top, width, height, *, word_wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    return box, tf


def write_lines(tf, lines, *, size_pt=18, bold=False, color=TEXT, font=BODY_FONT, align=PP_ALIGN.LEFT, space_after=6):
    first = True
    for item in lines:
        if isinstance(item, tuple):
            text, opts = item[0], item[1]
        else:
            text, opts = item, {}
        p = _clear_paragraphs(tf) if first else tf.add_paragraph()
        first = False
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", space_after))
        run = p.add_run()
        run.text = text
        _set_run(
            run,
            size_pt=opts.get("size", size_pt),
            bold=opts.get("bold", bold),
            color=opts.get("color", color),
            font=opts.get("font", font),
        )


def paint_background(slide, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD_DARK
    bar.line.fill.background()
    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), height - Inches(0.08), width, Inches(0.08)
    )
    foot.fill.solid()
    foot.fill.fore_color.rgb = GOLD_DARK
    foot.line.fill.background()


def add_footer(slide, width, height, text=FOOTER):
    _, tf = add_textbox(
        slide,
        Inches(0.5),
        height - Inches(0.45),
        width - Inches(1.0),
        Inches(0.3),
    )
    write_lines(tf, [text], size_pt=10, color=TEXT_DIM, font=TITLE_FONT, align=PP_ALIGN.LEFT)


def add_num_badge(slide, number: str):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(0.28), Inches(0.7), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLOOD
    badge.line.color.rgb = GOLD
    badge.line.width = Pt(1)
    tf = badge.text_frame
    tf.word_wrap = False
    p = _clear_paragraphs(tf)
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = number
    _set_run(run, size_pt=14, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_title(slide, title: str, left=Inches(1.35), top=Inches(0.28), width=Inches(8.2)):
    _, tf = add_textbox(slide, left, top, width, Inches(0.55))
    write_lines(tf, [title], size_pt=22, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_panel(slide, left, top, width, height):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = GOLD_DARK
    panel.line.width = Pt(1.25)
    return panel


def add_slide_image(slide, path: Path, left, top, width, height):
    if not path.exists():
        print(f"AVISO: imagem ausente — {path}", file=sys.stderr)
        return None
    return slide.shapes.add_picture(str(path), left, top, width=width, height=height)


def write_code(tf, text: str, *, size_pt=11):
    """Um bloco de código monoespaçado (linhas via \\n)."""
    first = True
    for line in text.splitlines():
        p = _clear_paragraphs(tf) if first else tf.add_paragraph()
        first = False
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        p.line_spacing = 1.0
        run = p.add_run()
        run.text = line if line else " "
        _set_run(run, size_pt=size_pt, bold=False, color=TEXT, font=CODE_FONT)


def add_code_slide(prs, *, badge: str, title: str, code: str, size_pt: float = 9):
    """Slide de código com painel contido acima do rodapé (slide 10×5.625\")."""
    s = new_slide(prs)
    W, H = prs.slide_width, prs.slide_height
    add_num_badge(s, badge)
    add_title(s, title)
    # Rodapé ocupa ~0.5"; painel termina em ~4.95"
    add_panel(s, Inches(0.4), Inches(0.88), Inches(9.2), Inches(4.0))
    _, tf = add_textbox(s, Inches(0.55), Inches(0.98), Inches(8.9), Inches(3.75))
    write_code(tf, code, size_pt=size_pt)
    add_footer(s, W, H)
    return s


def delete_all_slides(prs: Presentation):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def new_slide(prs: Presentation):
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    for shape in list(slide.shapes):
        sp = shape._element
        sp.getparent().remove(sp)
    paint_background(slide, prs.slide_width, prs.slide_height)
    return slide


def build_deck() -> Presentation:
    if not SRC.exists():
        raise SystemExit(f"Template não encontrado: {SRC}")

    prs = Presentation(str(SRC))
    delete_all_slides(prs)
    W, H = prs.slide_width, prs.slide_height

    # 1 — Capa
    s = new_slide(prs)
    add_num_badge(s, "06")
    _, tf = add_textbox(s, Inches(0.7), Inches(1.25), Inches(8.6), Inches(2.1))
    write_lines(
        tf,
        [
            ("Aula 06: Mercado, PI e", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 6}),
            ("Monetização Ética", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 16}),
            ("Mercado BR/internacional · Loja cosmética com moedas ganhas jogando", {"size": 14, "color": TEXT_DIM, "font": BODY_FONT}),
        ],
    )
    add_panel(s, Inches(0.7), Inches(3.7), Inches(4.6), Inches(0.9))
    _, tf = add_textbox(s, Inches(0.9), Inches(3.85), Inches(4.2), Inches(0.7))
    write_lines(
        tf,
        [
            ("MÓDULO 2", {"size": 12, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 2}),
            ("INTRODUÇÃO AO GDSCRIPT “DO ZERO”", {"size": 11, "color": TEXT, "font": BODY_FONT}),
        ],
    )
    add_footer(s, W, H)

    # 2 — Onde estamos
    s = new_slide(prs)
    add_num_badge(s, "01")
    add_title(s, "Onde estamos na trilha")
    steps = [
        ("PROVAÇÃO M1", "Círculo Mágico\nfechado"),
        ("AULA 06", "Mercado + loja\néticas (GDScript)"),
        ("AULA 07+", "Próximos passos\ndo script"),
    ]
    for i, (k, v) in enumerate(steps):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.4), Inches(3.0), Inches(2.5))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.65), Inches(2.6), Inches(2.1))
        lines = [
            (k, {"size": 14, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
        ]
        for part in v.split("\n"):
            lines.append((part, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 2}))
        write_lines(tf, lines)
        if i < 2:
            _, tf = add_textbox(s, left + Inches(2.7), Inches(2.4), Inches(0.6), Inches(0.4))
            write_lines(tf, [("→", {"size": 20, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER})])
    add_footer(s, W, H)

    # 3 — Ementa
    s = new_slide(prs)
    add_num_badge(s, "02")
    add_title(s, "Tópico da ementa")
    add_panel(s, Inches(0.55), Inches(1.4), Inches(8.9), Inches(2.8))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.9), Inches(8.2), Inches(2.0))
    write_lines(
        tf,
        [
            ("Mercado brasileiro e", {"size": 22, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("internacional de jogos", {"size": 22, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Receita ética · Original IP vs serviços · loja sem azar na Godot.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 4 — Como o dinheiro circula
    s = new_slide(prs)
    add_num_badge(s, "03")
    add_title(s, "Como o dinheiro circula")
    cards = [
        ("PREMIUM", "Preço único justo — você compra o jogo."),
        ("DLC / PASS", "Conteúdo claro e transparente."),
        ("COSMÉTICOS", "Skin sem vantagem injusta."),
        ("SERVIÇOS", "Arte, código, gamificação sob contrato."),
    ]
    for i, (k, v) in enumerate(cards):
        col = i % 2
        row = i // 2
        left = Inches(0.55 + col * 4.6)
        top = Inches(1.15 + row * 1.7)
        add_panel(s, left, top, Inches(4.35), Inches(1.5))
        _, tf = add_textbox(s, left + Inches(0.25), top + Inches(0.25), Inches(3.9), Inches(1.1))
        write_lines(
            tf,
            [
                (k, {"size": 14, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 8}),
                (v, {"size": 13, "color": TEXT}),
            ],
        )
    add_footer(s, W, H)

    # 5 — Original IP
    s = new_slide(prs)
    add_num_badge(s, "04")
    add_title(s, "Original IP")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.55), Inches(8.2), Inches(2.7))
    write_lines(
        tf,
        [
            ("Você cria e detém a propriedade intelectual.", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Jogo autoral · personagens · mundo · merch · sequências.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("O ativo é seu — e o contrato com o jogador também.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 6 — Prestação de serviços
    s = new_slide(prs)
    add_num_badge(s, "05")
    add_title(s, "Prestação de serviços")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.55), Inches(8.2), Inches(2.7))
    write_lines(
        tf,
        [
            ("Você vende capacidade — não necessariamente a IP.", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Art outsourcing · serious games · apps gamificados.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Ética e contrato mudam: quem é dono? O que vai no portfólio?", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 7 — Monetização ética
    s = new_slide(prs)
    add_num_badge(s, "06")
    add_title(s, "Monetização ética")
    points = [
        "Preço justo e transparente.",
        "DLC / pass com conteúdo claro.",
        "Cosméticos sem pay-to-win.",
        "Moeda ganha jogando — recompensa tempo e habilidade.",
    ]
    for i, text in enumerate(points):
        top = Inches(1.15 + i * 0.7)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.58))
        _, tf = add_textbox(s, Inches(0.85), top + Inches(0.1), Inches(8.4), Inches(0.4))
        write_lines(tf, [(f"◆  {text}", {"size": 15, "color": TEXT})])
    add_footer(s, W, H)

    # 8 — Microtx aceitável vs abusiva
    s = new_slide(prs)
    add_num_badge(s, "07")
    add_title(s, "Microtransações: aceitável × abusivo")
    add_panel(s, Inches(0.55), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(0.8), Inches(1.45), Inches(3.8), Inches(2.7))
    write_lines(
        tf,
        [
            ("ACEITÁVEL", {"size": 16, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Skin com moedas da fase", {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("Preço fixo visível", {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("Recompensa por jogar", {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_panel(s, Inches(5.15), Inches(1.2), Inches(4.3), Inches(3.1))
    _, tf = add_textbox(s, Inches(5.4), Inches(1.45), Inches(3.8), Inches(2.7))
    write_lines(
        tf,
        [
            ("ABUSIVO", {"size": 16, "bold": True, "color": BLOOD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Loot box com chance oculta", {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("Pay-to-win", {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("Dark patterns / pressão em crianças", {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 9 — Loot boxes + ClassInd
    s = new_slide(prs)
    add_num_badge(s, "08")
    add_title(s, "Loot boxes + ClassInd")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.55), Inches(8.2), Inches(2.7))
    write_lines(
        tf,
        [
            ("No Brasil, simular jogo de azar eleva a faixa.", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Loot boxes pagas → frequentemente 18+ (eco Aula 05).", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Hoje: preço fixo · efeito determinístico · sem randi() na compra.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 10 — Ponte prática
    s = new_slide(prs)
    add_num_badge(s, "09")
    add_title(s, "Ponte prática")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.7), Inches(8.2), Inches(2.5))
    write_lines(
        tf,
        [
            ("Hoje montamos o balcão sem azar.", {"size": 20, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 16}),
            ("Cena Control · PanelContainer · Button · Label", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Primeiro GDScript de UI do Módulo 2.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 11 — Nós Control
    s = new_slide(prs)
    add_num_badge(s, "10")
    add_title(s, "Nós Control")
    cards = [
        ("CONTROL", "Família de UI — âncoras, não metros do mundo."),
        ("PANELCONTAINER", "Fundo da loja e de cada item (Aura/Chapeu/Capa)."),
        ("BUTTON / LABEL", "Clique e texto — o balcão fala."),
    ]
    for i, (k, v) in enumerate(cards):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.3), Inches(3.0), Inches(2.8))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.6), Inches(2.6), Inches(2.3))
        write_lines(
            tf,
            [
                (k, {"size": 13, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
                (v, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
    add_footer(s, W, H)

    # 12 — Árvore Scene (foto Godot)
    s = new_slide(prs)
    add_num_badge(s, "11")
    add_title(s, "Árvore Scene (Godot)")
    add_slide_image(s, HIERARQUIA_IMG, Inches(0.45), Inches(1.05), Inches(3.55), Inches(4.0))
    add_panel(s, Inches(4.2), Inches(1.1), Inches(5.25), Inches(3.85))
    _, tf = add_textbox(s, Inches(4.4), Inches(1.25), Inches(4.9), Inches(3.55))
    write_lines(
        tf,
        [
            ("Loja (Control) ← loja.gd", {"size": 12, "bold": True, "color": GOLD_BRIGHT, "font": CODE_FONT, "space_after": 4}),
            ("└─ PanelContainer", {"size": 11, "color": TEXT, "font": CODE_FONT, "space_after": 2}),
            ("   └─ MarginContainer", {"size": 11, "color": TEXT, "font": CODE_FONT, "space_after": 2}),
            ("      └─ VBoxContainer", {"size": 11, "color": TEXT, "font": CODE_FONT, "space_after": 2}),
            ("         ├─ Título (Label)", {"size": 11, "color": TEXT, "font": CODE_FONT, "space_after": 2}),
            ("         ├─ Moedas (Label)", {"size": 11, "color": TEXT, "font": CODE_FONT, "space_after": 2}),
            ("         ├─ Aura (PanelContainer)", {"size": 11, "color": TEXT, "font": CODE_FONT, "space_after": 2}),
            ("         │    └─ HBox → Label + Button", {"size": 11, "color": TEXT, "font": CODE_FONT, "space_after": 2}),
            ("         ├─ Chapeu / Capa (idem)", {"size": 11, "color": TEXT, "font": CODE_FONT, "space_after": 2}),
            ("         ├─ GanharMoedas (Button)", {"size": 11, "color": TEXT, "font": CODE_FONT, "space_after": 2}),
            ("         └─ Status (Label)", {"size": 11, "color": TEXT, "font": CODE_FONT, "space_after": 8}),
            ("Margin = padding 10–16 · HBox: Label Expand + Button", {"size": 11, "color": TEXT_DIM}),
        ],
    )
    add_footer(s, W, H)

    # 13 — Sinais
    s = new_slide(prs)
    add_num_badge(s, "12")
    add_title(s, "Sinais: pressed → função")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.55), Inches(8.2), Inches(2.7))
    write_lines(
        tf,
        [
            ("O botão avisa. O script responde.", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("Aba Sinais → pressed() → Connect no nó Loja.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("GanharMoedas → _on_btn_ganhar_moedas_pressed", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("Button (Chapeu/Capa/Aura) → _on_btn_comprar_*_pressed", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 14 — Economia in-game
    s = new_slide(prs)
    add_num_badge(s, "13")
    add_title(s, "Economia in-game")
    cards = [
        ("SALDO 10", "Moedas já “coletadas” na fase."),
        ("+5 / COLETA", "Botão simula ganhar jogando."),
        ("5 · 12 · 20", "Chapéu · Capa · Aura — preço fixo."),
    ]
    for i, (k, v) in enumerate(cards):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.3), Inches(3.0), Inches(2.8))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.6), Inches(2.6), Inches(2.3))
        write_lines(
            tf,
            [
                (k, {"size": 16, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
                (v, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
    add_footer(s, W, H)

    # 15 — Código bloco 3
    add_code_slide(
        prs,
        badge="14",
        title="Script — Bloco 3 (base + coletar)",
        size_pt=9,
        code="""extends Control
var moedas: int = 10
var tem_chapeu: bool = false
var tem_capa: bool = false
var tem_aura: bool = false
const PRECO_CHAPEU: int = 5
const PRECO_CAPA: int = 12
const PRECO_AURA: int = 20
const GANHO_FASE: int = 5
@onready var label_moedas: Label = $PanelContainer/MarginContainer/VBoxContainer/Moedas
@onready var label_status: Label = $PanelContainer/MarginContainer/VBoxContainer/Status
@onready var btn_chapeu: Button = $PanelContainer/MarginContainer/VBoxContainer/Chapeu/HBoxContainer/Button
@onready var btn_capa: Button = $PanelContainer/MarginContainer/VBoxContainer/Capa/HBoxContainer/Button
@onready var btn_aura: Button = $PanelContainer/MarginContainer/VBoxContainer/Aura/HBoxContainer/Button
func _ready() -> void:
	_atualizar_hud()
	label_status.text = "Bem-vindo. Só moedas ganhas jogando."
func _atualizar_hud() -> void:
	label_moedas.text = "Moedas: %d" % moedas
func _on_btn_ganhar_moedas_pressed() -> void:
	moedas += GANHO_FASE
	label_status.text = "Você coletou +%d moedas na fase." % GANHO_FASE
	_atualizar_hud()""",
    )

    # 16 — Código bloco 4
    add_code_slide(
        prs,
        badge="15",
        title="Script — Bloco 4 (compra ética)",
        size_pt=11,
        code="""func _on_btn_comprar_chapeu_pressed() -> void:
	if tem_chapeu:
		label_status.text = "Você já tem chapéu."
		return
	if moedas < PRECO_CHAPEU:
		label_status.text = "Moedas insuficientes."
		return
	moedas -= PRECO_CHAPEU
	tem_chapeu = true
	btn_chapeu.disabled = true
	btn_chapeu.text = "Adquirido"
	label_status.text = "Chapéu comprado!"
	_atualizar_hud()

# Repita o mesmo padrão para Capa e Aura
# (PRECO_CAPA / PRECO_AURA · tem_capa / tem_aura · btn_capa / btn_aura)

# Proibido: randi() · loot box · preço em dinheiro real""",
    )

    # 17 — Checklist
    s = new_slide(prs)
    add_num_badge(s, "16")
    add_title(s, "Checklist do artefato")
    checks = [
        "ui/loja.tscn com raiz Control nomeada Loja.",
        "PanelContainer + 3 itens (Label+Button) + GanharMoedas.",
        "loja.gd anexado · sinal pressed conectado.",
        "Label Moedas atualiza · bloqueios sem saldo / já possui.",
        "Sem randi() / loot / IAP real.",
        "F6 demonstra o fluxo · anotações enviadas.",
    ]
    for i, text in enumerate(checks):
        top = Inches(1.05 + i * 0.52)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.45))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.05), Inches(8.4), Inches(0.35))
        write_lines(tf, [(f"□  {text}", {"size": 12, "color": TEXT})])
    add_footer(s, W, H)

    # 18 — Fechamento
    s = new_slide(prs)
    add_num_badge(s, "17")
    add_title(s, "Fechamento e Altar")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.55), Inches(8.2), Inches(2.7))
    write_lines(
        tf,
        [
            ("FIM DA AULA 06: GUARDIÃO DA LOJA ÉTICA", {"size": 15, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Respeitar o tempo do jogador também é modelo de negócio.", {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Envie as anotações. Leve o código ao Altar quando o Mestre liberar.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Próxima trilha: continuar o GDScript “do zero” (Aula 07).", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    return prs


def export_pdf_with_powerpoint(pptx_path: Path, pdf_path: Path) -> bool:
    """Exporta PDF via Microsoft PowerPoint (Windows COM)."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        ps = f"""
$ppt = New-Object -ComObject PowerPoint.Application
$ppt.Visible = [Microsoft.Office.Core.MsoTriState]::msoTrue
$pres = $ppt.Presentations.Open('{pptx_path}', $true, $false, $false)
$pres.SaveAs('{pdf_path}', 32)
$pres.Close()
$ppt.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($pres) | Out-Null
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null
"""
        import subprocess

        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            print(result.stdout)
            print(result.stderr, file=sys.stderr)
            return False
        return pdf_path.exists()

    powerpoint = win32com.client.Dispatch("PowerPoint.Application")
    powerpoint.Visible = 1
    presentation = powerpoint.Presentations.Open(str(pptx_path), WithWindow=False)
    presentation.SaveAs(str(pdf_path), 32)
    presentation.Close()
    powerpoint.Quit()
    return pdf_path.exists()


def main():
    prs = build_deck()
    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"PPTX: {OUT_PPTX} ({OUT_PPTX.stat().st_size} bytes)")

    if OUT_PDF.exists():
        OUT_PDF.unlink()
    ok = export_pdf_with_powerpoint(OUT_PPTX.resolve(), OUT_PDF.resolve())
    if ok:
        print(f"PDF:  {OUT_PDF} ({OUT_PDF.stat().st_size} bytes)")
    else:
        print("PDF:  não gerado (PowerPoint COM indisponível). PPTX ok.", file=sys.stderr)
        sys.exit(0)


if __name__ == "__main__":
    main()
