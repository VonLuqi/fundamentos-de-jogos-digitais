#!/usr/bin/env python3
"""Gera os slides da Aula 07 a partir do template visual da Aula 01.

Saídas:
  assets/docs/aulas/aula07_papeis_workflow_slides.pptx
  assets/docs/aulas/aula07_papeis_workflow_slides.pdf  (via PowerPoint COM, se disponível)
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pptx.oxml import parse_xml

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "assets" / "docs" / "aulas" / "aula01_godot_slides.pptx"
OUT_PPTX = ROOT / "assets" / "docs" / "aulas" / "aula07_papeis_workflow_slides.pptx"
OUT_PDF = ROOT / "assets" / "docs" / "aulas" / "aula07_papeis_workflow_slides.pdf"
IMG_DIR = ROOT / "assets" / "docs" / "aulas" / "aula07-equipe-labirinto"
IMG_HYTALE = IMG_DIR / "hytale-scope.jpg"
IMG_RIOT = IMG_DIR / "riot-games-logo.png"

BG = RGBColor(0x0D, 0x0A, 0x10)
PANEL = RGBColor(0x18, 0x12, 0x1A)
GOLD = RGBColor(0xCF, 0xA7, 0x59)
GOLD_BRIGHT = RGBColor(0xF2, 0xD5, 0x9A)
GOLD_DARK = RGBColor(0x7A, 0x5C, 0x2E)
TEXT = RGBColor(0xEC, 0xE1, 0xD1)
TEXT_DIM = RGBColor(0xAB, 0x9C, 0x8A)
BLOOD = RGBColor(0x7A, 0x1F, 0x2B)

TITLE_FONT = "Cinzel"
BODY_FONT = "Georgia"
FOOTER = "AULA 07 | MÓDULO 2"


def _set_run(run, *, size_pt, bold=False, color=TEXT, font=BODY_FONT):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "cs", "ea"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = parse_xml(
                f'<a:{tag} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'
            )
            rPr.append(el)
        else:
            el.set("typeface", font)


def _clear_paragraphs(tf):
    p = tf.paragraphs[0]
    p.clear()
    for extra in list(tf.paragraphs)[1:]:
        p_elem = extra._p
        p_elem.getparent().remove(p_elem)
    return tf.paragraphs[0]


def add_textbox(slide, left, top, width, height, *, word_wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    return box, tf


def write_lines(tf, lines, *, size_pt=18, bold=False, color=TEXT, font=BODY_FONT, align=PP_ALIGN.LEFT, space_after=6):
    first = True
    for item in lines:
        if isinstance(item, tuple):
            text, opts = item[0], item[1]
        else:
            text, opts = item, {}
        p = _clear_paragraphs(tf) if first else tf.add_paragraph()
        first = False
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", space_after))
        run = p.add_run()
        run.text = text
        _set_run(
            run,
            size_pt=opts.get("size", size_pt),
            bold=opts.get("bold", bold),
            color=opts.get("color", color),
            font=opts.get("font", font),
        )


def paint_background(slide, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD_DARK
    bar.line.fill.background()
    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), height - Inches(0.08), width, Inches(0.08)
    )
    foot.fill.solid()
    foot.fill.fore_color.rgb = GOLD_DARK
    foot.line.fill.background()


def add_footer(slide, width, height, text=FOOTER):
    _, tf = add_textbox(
        slide,
        Inches(0.5),
        height - Inches(0.45),
        width - Inches(1.0),
        Inches(0.3),
    )
    write_lines(tf, [text], size_pt=10, color=TEXT_DIM, font=TITLE_FONT, align=PP_ALIGN.LEFT)


def add_num_badge(slide, number: str):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(0.28), Inches(0.7), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLOOD
    badge.line.color.rgb = GOLD
    badge.line.width = Pt(1)
    tf = badge.text_frame
    tf.word_wrap = False
    p = _clear_paragraphs(tf)
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = number
    _set_run(run, size_pt=14, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_title(slide, title: str, left=Inches(1.35), top=Inches(0.28), width=Inches(8.2)):
    _, tf = add_textbox(slide, left, top, width, Inches(0.55))
    write_lines(tf, [title], size_pt=22, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_panel(slide, left, top, width, height):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = GOLD_DARK
    panel.line.width = Pt(1.25)
    return panel


def add_slide_image(slide, path: Path, left, top, width=None, height=None):
    if not path.exists():
        print(f"AVISO: imagem ausente — {path}", file=sys.stderr)
        return None
    kwargs = {}
    if width is not None:
        kwargs["width"] = width
    if height is not None:
        kwargs["height"] = height
    return slide.shapes.add_picture(str(path), left, top, **kwargs)


def delete_all_slides(prs: Presentation):
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def new_slide(prs: Presentation):
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    for shape in list(slide.shapes):
        sp = shape._element
        sp.getparent().remove(sp)
    paint_background(slide, prs.slide_width, prs.slide_height)
    return slide


def build_deck() -> Presentation:
    if not SRC.exists():
        raise SystemExit(f"Template não encontrado: {SRC}")

    prs = Presentation(str(SRC))
    delete_all_slides(prs)
    W, H = prs.slide_width, prs.slide_height

    # 1 — Capa
    s = new_slide(prs)
    add_num_badge(s, "07")
    _, tf = add_textbox(s, Inches(0.7), Inches(1.25), Inches(8.6), Inches(2.1))
    write_lines(
        tf,
        [
            ("Aula 07: Papéis, Workflow e", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 6}),
            ("Versionamento Visual", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 16}),
            ("Individual ou dupla · arte da moeda × coleta · diário compartilhado", {"size": 14, "color": TEXT_DIM, "font": BODY_FONT}),
        ],
    )
    add_panel(s, Inches(0.7), Inches(3.7), Inches(4.6), Inches(0.9))
    _, tf = add_textbox(s, Inches(0.9), Inches(3.85), Inches(4.2), Inches(0.7))
    write_lines(
        tf,
        [
            ("MÓDULO 2", {"size": 12, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 2}),
            ("INTRODUÇÃO AO GDSCRIPT “DO ZERO”", {"size": 11, "color": TEXT, "font": BODY_FONT}),
        ],
    )
    add_footer(s, W, H)

    # 2 — Onde estamos
    s = new_slide(prs)
    add_num_badge(s, "01")
    add_title(s, "Onde estamos na trilha")
    steps = [
        ("AULA 06", "Loja ética\n(2 encontros)"),
        ("AULA 07", "Papéis + coleta\n(2 encontros)"),
        ("AULA 08+", "Expandir o\nLabirinto"),
    ]
    for i, (k, v) in enumerate(steps):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.4), Inches(3.0), Inches(2.5))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.65), Inches(2.6), Inches(2.1))
        lines = [
            (k, {"size": 14, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
        ]
        for part in v.split("\n"):
            lines.append((part, {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 2}))
        write_lines(tf, lines)
        if i < 2:
            _, tf = add_textbox(s, left + Inches(2.7), Inches(2.4), Inches(0.6), Inches(0.4))
            write_lines(tf, [("→", {"size": 20, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER})])
    add_footer(s, W, H)

    # 3 — Ementa
    s = new_slide(prs)
    add_num_badge(s, "02")
    add_title(s, "Tópico da ementa")
    add_panel(s, Inches(0.55), Inches(1.4), Inches(8.9), Inches(2.8))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.7), Inches(8.2), Inches(2.3))
    write_lines(
        tf,
        [
            ("Papéis no desenvolvimento de jogos", {"size": 20, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("e Workflow de desenvolvimento", {"size": 20, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Scope Creep · individual ou dupla · arte da moeda × coleta · diário compartilhado.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 4 — Cinco ofícios (denso)
    s = new_slide(prs)
    add_num_badge(s, "03")
    add_title(s, "Cinco ofícios — entregável e falha típica")
    cards = [
        ("GAME DESIGN", "Entrega: regras / GDD\nFalha: “a gente vê\ndepois se é divertido”"),
        ("PROGRAMAÇÃO", "Entrega: cena que\nfunciona\nFalha: script sem\ndono da .tscn"),
        ("ARTE", "Entrega: sprite / cena\nvisível\nFalha: arte sem\nnome de arquivo"),
        ("ÁUDIO", "Entrega: SFX/música\nFalha: “deixamos\nsilêncio para depois”"),
        ("PRODUÇÃO", "Entrega: cercado +\nsync\nFalha: ninguém diz\nnão ao Scope Creep"),
    ]
    for i, (k, v) in enumerate(cards):
        left = Inches(0.3 + i * 1.96)
        add_panel(s, left, Inches(1.15), Inches(1.88), Inches(3.35))
        _, tf = add_textbox(s, left + Inches(0.08), Inches(1.3), Inches(1.72), Inches(3.05))
        lines = [
            (k, {"size": 10, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 8}),
        ]
        for part in v.split("\n"):
            lines.append((part, {"size": 10, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 2}))
        write_lines(tf, lines)
    add_footer(s, W, H)

    # 5 — Workflow denso
    s = new_slide(prs)
    add_num_badge(s, "04")
    add_title(s, "Workflow — o que entra e sai")
    steps = [
        ("IDEIA", "Entra: sonho\nSai: intenção\nclara"),
        ("MVP", "Entra: lista\nSai: cercado\nfechado"),
        ("PRODUÇÃO", "Entra: ofício\nSai: arquivo\ncom dono"),
        ("INTEGRAÇÃO", "Entra: cenas\nSai: um projeto\nque abre"),
        ("PLAYTEST", "Entra: build\nSai: corte ou\npróximo passo"),
    ]
    for i, (k, v) in enumerate(steps):
        left = Inches(0.3 + i * 1.96)
        add_panel(s, left, Inches(1.25), Inches(1.88), Inches(3.1))
        _, tf = add_textbox(s, left + Inches(0.08), Inches(1.45), Inches(1.72), Inches(2.8))
        lines = [
            (k, {"size": 11, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 8}),
        ]
        for part in v.split("\n"):
            lines.append((part, {"size": 11, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 2}))
        write_lines(tf, lines)
    add_footer(s, W, H)

    # 6 — Scope Creep denso
    s = new_slide(prs)
    add_num_badge(s, "05")
    add_title(s, "Scope Creep no Labirinto")
    cards = [
        ("EXEMPLO RUIM", "“E se tiver inimigos, loja, boss e online?” — ideias boas no momento errado."),
        ("MVP DESTA AULA", "E1: sprite LibreSprite · E2: Player + coleta. Só isso."),
        ("ANTÍDOTO", "Lista FORA DO ESCOPO no diário · dizer não · loja Aula 06 fica fora."),
    ]
    for i, (k, v) in enumerate(cards):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.3), Inches(3.0), Inches(2.8))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.55), Inches(2.6), Inches(2.4))
        write_lines(
            tf,
            [
                (k, {"size": 12, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
                (v, {"size": 12, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
    add_footer(s, W, H)

    # 7 — Caso Hytale (imagem + linha do tempo)
    s = new_slide(prs)
    add_num_badge(s, "06")
    add_title(s, "Caso: Hytale — visão enorme")
    add_slide_image(s, IMG_HYTALE, Inches(0.45), Inches(1.1), width=Inches(5.15))
    add_panel(s, Inches(5.75), Inches(1.1), Inches(3.8), Inches(3.35))
    _, tf = add_textbox(s, Inches(5.95), Inches(1.25), Inches(3.4), Inches(3.05))
    write_lines(
        tf,
        [
            ("LINHA DO TEMPO", {"size": 12, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 10}),
            ("2018 — Trailer: sandbox + RPG + tools + multiplayer…", {"size": 11, "color": TEXT, "space_after": 6}),
            ("2020 — Riot Games compra a Hypixel Studios.", {"size": 11, "color": TEXT, "space_after": 6}),
            ("2021+ — Beta não chega; reboot de engine; ambição cresce.", {"size": 11, "color": TEXT, "space_after": 6}),
            ("2025 — Projeto cancelado (depois reaberto com escopo cortado).", {"size": 11, "color": TEXT, "space_after": 10}),
            ("Dinheiro e talento ≠ cercado fechado.", {"size": 11, "bold": True, "color": GOLD, "space_after": 0}),
        ],
    )
    add_footer(s, W, H)

    # 8 — Riot × lição para a sala
    s = new_slide(prs)
    add_num_badge(s, "07")
    add_title(s, "Riot × Hytale — lição de escopo")
    add_panel(s, Inches(0.55), Inches(1.15), Inches(3.2), Inches(3.2))
    add_slide_image(s, IMG_RIOT, Inches(0.95), Inches(1.85), width=Inches(2.4))
    add_panel(s, Inches(3.95), Inches(1.15), Inches(5.5), Inches(3.2))
    _, tf = add_textbox(s, Inches(4.2), Inches(1.35), Inches(5.05), Inches(2.85))
    write_lines(
        tf,
        [
            ("O QUE ACONTECEU", {"size": 12, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 8}),
            ("A visão era tão grande que cada “só mais isso” empurrava o jogo para longe do lançável.", {"size": 12, "color": TEXT, "space_after": 8}),
            ("Até com backing da Riot, falta de MVP claro virou anos sem build pública estável.", {"size": 12, "color": TEXT, "space_after": 10}),
            ("NA NOSSA SALA", {"size": 12, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 8}),
            ("Labirinto = sprite + coleta. Loja, boss e online ficam FORA — senão viramos Hytale de 120 minutos.", {"size": 12, "color": TEXT, "space_after": 0}),
        ],
    )
    add_footer(s, W, H)

    # 9 — Versionamento visual
    s = new_slide(prs)
    add_num_badge(s, "08")
    add_title(s, "Versionamento visual")
    points = [
        "Nome canônico de pasta/cena = contrato (moeda.tscn ≠ MoedaFinal2).",
        "Projeto Godot só na máquina local — sem pasta sync / Drive.",
        "PNG do LibreSprite entra em sprites/ no Encontro 2.",
        "Sem dono de arquivo, alguém apaga o trabalho do outro sem perceber.",
    ]
    for i, text in enumerate(points):
        top = Inches(1.15 + i * 0.7)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.58))
        _, tf = add_textbox(s, Inches(0.85), top + Inches(0.1), Inches(8.4), Inches(0.4))
        write_lines(tf, [(f"◆  {text}", {"size": 14, "color": TEXT})])
    add_footer(s, W, H)

    # 10 — Ponte prática
    s = new_slide(prs)
    add_num_badge(s, "09")
    add_title(s, "Ponte prática — 2 encontros")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.55), Inches(8.2), Inches(2.7))
    write_lines(
        tf,
        [
            ("Oficina em dois encontros · individual ou dupla", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
            ("E1: teoria · convite · arte no LibreSprite", {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("E2: Godot local · coleta · diário · Altar", {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 11 — Estúdio → sala (abertura da prática)
    s = new_slide(prs)
    add_num_badge(s, "10")
    add_title(s, "Estúdio → sala (dupla ou solo)")
    add_panel(s, Inches(0.55), Inches(1.2), Inches(4.3), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.8), Inches(1.4), Inches(3.8), Inches(2.85))
    write_lines(
        tf,
        [
            ("DUPLA", {"size": 16, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Arte da moeda ↔ LibreSprite + moeda.tscn", {"size": 12, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("Programação ↔ Player + coleta", {"size": 12, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("Convite na própria Oficina.", {"size": 12, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_panel(s, Inches(5.15), Inches(1.2), Inches(4.3), Inches(3.15))
    _, tf = add_textbox(s, Inches(5.4), Inches(1.4), Inches(3.8), Inches(2.85))
    write_lines(
        tf,
        [
            ("INDIVIDUAL", {"size": 16, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Você faz os dois ofícios", {"size": 12, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("(arte no E1 · Godot no E2).", {"size": 12, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 8}),
            ("Só 1 ou 2 pessoas — sem trio.", {"size": 12, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 12 — Ofícios da sala
    s = new_slide(prs)
    add_num_badge(s, "11")
    add_title(s, "Ofícios da sala")
    cards = [
        ("ARTE DA MOEDA", "E1: LibreSprite\nE2: moeda.tscn\n+ Area2D"),
        ("PROGRAMAÇÃO", "E2: Player (M1)\nbody_entered\ncoleta / conta 1"),
        ("SOLO", "Os dois ofícios\narte no E1\nGodot no E2"),
    ]
    for i, (k, v) in enumerate(cards):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.3), Inches(3.0), Inches(2.8))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.5), Inches(2.6), Inches(2.5))
        lines = [
            (k, {"size": 13, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 10}),
        ]
        for part in v.split("\n"):
            lines.append((part, {"size": 12, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 2}))
        write_lines(tf, lines)
    add_footer(s, W, H)

    # 13 — Diário
    s = new_slide(prs)
    add_num_badge(s, "12")
    add_title(s, "Diário compartilhado")
    points = [
        "Substitui anotações + síntese — um texto só.",
        "Em dupla: os dois editam o mesmo diário (autosave + sync).",
        "Seções: ofícios · entregas · cercado · versionamento.",
        "Cada aluno finaliza o próprio envio no Encontro 2.",
    ]
    for i, text in enumerate(points):
        top = Inches(1.15 + i * 0.7)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.58))
        _, tf = add_textbox(s, Inches(0.85), top + Inches(0.1), Inches(8.4), Inches(0.4))
        write_lines(tf, [(f"◆  {text}", {"size": 15, "color": TEXT})])
    add_footer(s, W, H)

    # 14 — Pastas Godot
    s = new_slide(prs)
    add_num_badge(s, "13")
    add_title(s, "Pastas Godot (contrato)")
    add_panel(s, Inches(0.55), Inches(1.15), Inches(8.9), Inches(3.25))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.35), Inches(8.2), Inches(2.9))
    write_lines(
        tf,
        [
            ("LabirintoDeMoedas/", {"size": 14, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 6}),
            ("cenas/  player.tscn · moeda.tscn · cenario.tscn (stub)", {"size": 13, "color": TEXT, "space_after": 4}),
            ("sprites/ · audio/ · scripts/ · ui/", {"size": 13, "color": TEXT, "space_after": 10}),
            ("MVP: sprite + movimento M1 + coleta (Area2D).", {"size": 12, "color": TEXT_DIM}),
        ],
    )
    add_footer(s, W, H)

    # 15 — Projeto local
    s = new_slide(prs)
    add_num_badge(s, "14")
    add_title(s, "Projeto local (sem pasta sync)")
    checks = [
        "Godot só na máquina local — sem Drive/OneDrive no projeto.",
        "PNG do LibreSprite → sprites/ no Encontro 2.",
        "Na dupla: cada um na cena do seu ofício.",
        "Nomes canônicos: player.tscn · moeda.tscn.",
        "Pasta .godot/ é cache local — não versionar.",
    ]
    for i, text in enumerate(checks):
        top = Inches(1.1 + i * 0.55)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.48))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.05), Inches(8.4), Inches(0.38))
        write_lines(tf, [(f"◆  {text}", {"size": 13, "color": TEXT})])
    add_footer(s, W, H)

    # 16 — Checklist
    s = new_slide(prs)
    add_num_badge(s, "15")
    add_title(s, "Checklist do artefato")
    checks = [
        "Solo ou dupla com ofícios definidos na Oficina.",
        "E1: sprite LibreSprite · E2: moeda.tscn + coleta.",
        "≥3 itens fora do escopo no diário.",
        "Projeto Godot local (sem pasta sync).",
        "Diário finalizado (envio individual) + Altar.",
    ]
    for i, text in enumerate(checks):
        top = Inches(1.1 + i * 0.55)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.48))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.05), Inches(8.4), Inches(0.38))
        write_lines(tf, [(f"□  {text}", {"size": 13, "color": TEXT})])
    add_footer(s, W, H)

    # 17 — Fechamento
    s = new_slide(prs)
    add_num_badge(s, "16")
    add_title(s, "Fechamento e Altar")
    add_panel(s, Inches(0.55), Inches(1.25), Inches(8.9), Inches(3.15))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.55), Inches(8.2), Inches(2.7))
    write_lines(
        tf,
        [
            ("FIM DA AULA 07: CARTÓGRAFO DA DUPLA", {"size": 15, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Um labirinto pequeno e terminado ensina mais que um mundo aberto abandonado.", {"size": 13, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Finalize o diário. Leve o código ao Altar quando o Mestre liberar.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("Próxima trilha: expandir o Labirinto (Aula 08+).", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    return prs


def export_pdf_with_powerpoint(pptx_path: Path, pdf_path: Path) -> bool:
    """Exporta PDF via Microsoft PowerPoint (Windows COM)."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        ps = f"""
$ppt = New-Object -ComObject PowerPoint.Application
$ppt.Visible = [Microsoft.Office.Core.MsoTriState]::msoTrue
$pres = $ppt.Presentations.Open('{pptx_path}', $true, $false, $false)
$pres.SaveAs('{pdf_path}', 32)
$pres.Close()
$ppt.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($pres) | Out-Null
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null
"""
        import subprocess

        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            print(result.stdout)
            print(result.stderr, file=sys.stderr)
            return False
        return pdf_path.exists()

    powerpoint = win32com.client.Dispatch("PowerPoint.Application")
    powerpoint.Visible = 1
    presentation = powerpoint.Presentations.Open(str(pptx_path), WithWindow=False)
    presentation.SaveAs(str(pdf_path), 32)
    presentation.Close()
    powerpoint.Quit()
    return pdf_path.exists()


def main():
    prs = build_deck()
    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"PPTX: {OUT_PPTX} ({OUT_PPTX.stat().st_size} bytes)")

    if OUT_PDF.exists():
        OUT_PDF.unlink()
    ok = export_pdf_with_powerpoint(OUT_PPTX.resolve(), OUT_PDF.resolve())
    if ok:
        print(f"PDF:  {OUT_PDF} ({OUT_PDF.stat().st_size} bytes)")
    else:
        print("PDF:  não gerado (PowerPoint COM indisponível). PPTX ok.", file=sys.stderr)
        sys.exit(0)


if __name__ == "__main__":
    main()
