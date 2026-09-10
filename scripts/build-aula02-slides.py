#!/usr/bin/env python3
"""Gera os slides da Aula 02 a partir do template visual da Aula 01.

Saídas:
  assets/docs/aulas/aula02_glossario_player_slides.pptx
  assets/docs/aulas/aula02_glossario_player_slides.pdf  (via PowerPoint COM, se disponível)
"""

from __future__ import annotations

import sys
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from pptx.oxml import parse_xml

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "assets" / "docs" / "aulas" / "aula01_godot_slides.pptx"
OUT_PPTX = ROOT / "assets" / "docs" / "aulas" / "aula02_glossario_player_slides.pptx"
OUT_PDF = ROOT / "assets" / "docs" / "aulas" / "aula02_glossario_player_slides.pdf"

# Tokens Hades (css/hades-tokens.css)
BG = RGBColor(0x0D, 0x0A, 0x10)
PANEL = RGBColor(0x18, 0x12, 0x1A)
GOLD = RGBColor(0xCF, 0xA7, 0x59)
GOLD_BRIGHT = RGBColor(0xF2, 0xD5, 0x9A)
GOLD_DARK = RGBColor(0x7A, 0x5C, 0x2E)
TEXT = RGBColor(0xEC, 0xE1, 0xD1)
TEXT_DIM = RGBColor(0xAB, 0x9C, 0x8A)
BLOOD = RGBColor(0x7A, 0x1F, 0x2B)

TITLE_FONT = "Cinzel"
BODY_FONT = "Georgia"


def _set_run(run, *, size_pt, bold=False, color=TEXT, font=BODY_FONT):
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    # East Asian / latin hint helps PowerPoint pick the face
    rPr = run._r.get_or_add_rPr()
    for tag in ("latin", "cs", "ea"):
        el = rPr.find(qn(f"a:{tag}"))
        if el is None:
            el = parse_xml(
                f'<a:{tag} xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" typeface="{font}"/>'
            )
            rPr.append(el)
        else:
            el.set("typeface", font)


def _clear_paragraphs(tf):
    p = tf.paragraphs[0]
    p.clear()
    for extra in list(tf.paragraphs)[1:]:
        p_elem = extra._p
        p_elem.getparent().remove(p_elem)
    return tf.paragraphs[0]


def add_textbox(slide, left, top, width, height, *, word_wrap=True):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = word_wrap
    return box, tf


def write_lines(tf, lines, *, size_pt=18, bold=False, color=TEXT, font=BODY_FONT, align=PP_ALIGN.LEFT, space_after=6):
    """lines: list[str] or list[tuple]."""
    first = True
    for item in lines:
        if isinstance(item, tuple):
            text, opts = item[0], item[1]
        else:
            text, opts = item, {}
        p = _clear_paragraphs(tf) if first else tf.add_paragraph()
        first = False
        p.alignment = opts.get("align", align)
        p.space_after = Pt(opts.get("space_after", space_after))
        run = p.add_run()
        run.text = text
        _set_run(
            run,
            size_pt=opts.get("size", size_pt),
            bold=opts.get("bold", bold),
            color=opts.get("color", color),
            font=opts.get("font", font),
        )


def paint_background(slide, width, height):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = BG
    shape.line.fill.background()
    # top gold hairline
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(0), Emu(0), width, Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = GOLD_DARK
    bar.line.fill.background()
    # bottom accent
    foot = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Emu(0), height - Inches(0.08), width, Inches(0.08)
    )
    foot.fill.solid()
    foot.fill.fore_color.rgb = GOLD_DARK
    foot.line.fill.background()


def add_footer(slide, width, height, text="AULA 02 | MÓDULO 1"):
    _, tf = add_textbox(
        slide,
        Inches(0.5),
        height - Inches(0.45),
        width - Inches(1.0),
        Inches(0.3),
    )
    write_lines(tf, [text], size_pt=10, color=TEXT_DIM, font=TITLE_FONT, align=PP_ALIGN.LEFT)


def add_num_badge(slide, number: str):
    badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.45), Inches(0.28), Inches(0.7), Inches(0.42))
    badge.fill.solid()
    badge.fill.fore_color.rgb = BLOOD
    badge.line.color.rgb = GOLD
    badge.line.width = Pt(1)
    tf = badge.text_frame
    tf.word_wrap = False
    p = _clear_paragraphs(tf)
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = number
    _set_run(run, size_pt=14, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_title(slide, title: str, left=Inches(1.35), top=Inches(0.28), width=Inches(8.2)):
    _, tf = add_textbox(slide, left, top, width, Inches(0.55))
    write_lines(tf, [title], size_pt=22, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)


def add_panel(slide, left, top, width, height):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    panel.fill.solid()
    panel.fill.fore_color.rgb = PANEL
    panel.line.color.rgb = GOLD_DARK
    panel.line.width = Pt(1.25)
    return panel


def delete_all_slides(prs: Presentation):
    """Remove every slide while keeping the slide master/theme from the template."""
    sldIdLst = prs.slides._sldIdLst
    for sldId in list(sldIdLst):
        rId = sldId.get(qn("r:id"))
        prs.part.drop_rel(rId)
        sldIdLst.remove(sldId)


def new_slide(prs: Presentation):
    layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(layout)
    # Remove any leftover placeholders from DEFAULT layout
    for shape in list(slide.shapes):
        sp = shape._element
        sp.getparent().remove(sp)
    paint_background(slide, prs.slide_width, prs.slide_height)
    return slide


def build_deck() -> Presentation:
    if not SRC.exists():
        raise SystemExit(f"Template não encontrado: {SRC}")

    prs = Presentation(str(SRC))
    delete_all_slides(prs)
    W, H = prs.slide_width, prs.slide_height

    # 1 — Capa
    s = new_slide(prs)
    add_num_badge(s, "02")
    _, tf = add_textbox(s, Inches(0.7), Inches(1.6), Inches(8.6), Inches(1.4))
    write_lines(
        tf,
        [
            ("Aula 02: O Glossário do Desenvolvedor", {"size": 28, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 8}),
            ("e o Player na Tela", {"size": 28, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 16}),
            ("Core Loop, Grokking, Assets · Cenas, Nós, Input Map e movimento mínimo", {"size": 14, "color": TEXT_DIM, "font": BODY_FONT}),
        ],
    )
    add_panel(s, Inches(0.7), Inches(3.7), Inches(4.2), Inches(0.9))
    _, tf = add_textbox(s, Inches(0.9), Inches(3.85), Inches(3.8), Inches(0.7))
    write_lines(
        tf,
        [
            ("MÓDULO 1", {"size": 12, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 2}),
            ("FUNDAÇÕES, CULTURA E INTERFACE", {"size": 11, "color": TEXT, "font": BODY_FONT}),
        ],
    )
    add_footer(s, W, H)

    # 2 — Objetivos / artefato
    s = new_slide(prs)
    add_num_badge(s, "01")
    add_title(s, "Objetivos da aula e artefato")
    cards = [
        ("APRENDER", "Definir Core Loop, Grokking e Assets de forma operacional."),
        ("MONTAR", "Criar a cena do Player: CharacterBody2D + Sprite2D + CollisionShape2D."),
        ("MAPEAR", "Input Map com ir_cima, ir_baixo, ir_esquerda, ir_direita."),
        ("MOVER", "Script mínimo player.gd com move_and_slide no Play."),
    ]
    for i, (k, v) in enumerate(cards):
        col = i % 2
        row = i // 2
        left = Inches(0.55 + col * 4.6)
        top = Inches(1.1 + row * 1.7)
        add_panel(s, left, top, Inches(4.3), Inches(1.5))
        _, tf = add_textbox(s, left + Inches(0.2), top + Inches(0.2), Inches(3.9), Inches(1.15))
        write_lines(
            tf,
            [
                (k, {"size": 13, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 6}),
                (v, {"size": 14, "color": TEXT, "font": BODY_FONT}),
            ],
        )
    add_footer(s, W, H)

    # 3 — Glossário visão geral
    s = new_slide(prs)
    add_num_badge(s, "02")
    add_title(s, "Glossário — três termos do ofício")
    terms = [
        ("CORE LOOP", "O ciclo básico que o jogador repete."),
        ("GROKKING", "Controles na memória muscular."),
        ("ASSETS", "Imagens, sons e recursos do jogo."),
    ]
    for i, (k, v) in enumerate(terms):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.2), Inches(3.0), Inches(2.8))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.5), Inches(2.6), Inches(2.3))
        write_lines(
            tf,
            [
                (k, {"size": 14, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 14}),
                (v, {"size": 15, "color": TEXT, "font": BODY_FONT, "align": PP_ALIGN.CENTER}),
            ],
        )
    add_footer(s, W, H)

    # 4 — Core Loop
    s = new_slide(prs)
    add_num_badge(s, "03")
    add_title(s, "Core Loop — andar → coletar → avançar")
    _, tf = add_textbox(s, Inches(0.6), Inches(1.1), Inches(8.8), Inches(0.6))
    write_lines(
        tf,
        ["O ciclo básico que o jogador repete. Se o loop não for claro, o jogo não ensina o que fazer em seguida."],
        size_pt=15,
        color=TEXT_DIM,
    )
    steps = ["ANDAR", "COLETAR", "AVANÇAR"]
    for i, label in enumerate(steps):
        left = Inches(0.7 + i * 3.1)
        oval = s.shapes.add_shape(MSO_SHAPE.OVAL, left, Inches(2.1), Inches(2.2), Inches(2.2))
        oval.fill.solid()
        oval.fill.fore_color.rgb = PANEL
        oval.line.color.rgb = GOLD
        oval.line.width = Pt(2)
        _, tf = add_textbox(s, left, Inches(2.85), Inches(2.2), Inches(0.6))
        write_lines(tf, [(label, {"size": 16, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER})])
        if i < 2:
            _, tf = add_textbox(s, left + Inches(2.15), Inches(2.9), Inches(0.9), Inches(0.4))
            write_lines(tf, [("→", {"size": 22, "bold": True, "color": GOLD, "font": TITLE_FONT, "align": PP_ALIGN.CENTER})])
    add_footer(s, W, H)

    # 5 — Grokking
    s = new_slide(prs)
    add_num_badge(s, "04")
    add_title(s, "Grokking — memória muscular dos controles")
    add_panel(s, Inches(0.55), Inches(1.2), Inches(8.9), Inches(3.2))
    _, tf = add_textbox(s, Inches(0.85), Inches(1.5), Inches(8.3), Inches(2.7))
    write_lines(
        tf,
        [
            ("O momento em que os controles saem da cabeça e entram nas mãos.", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 14}),
            ("Antes do Grokking: o jogador pensa “preciso apertar W”.", {"size": 15, "color": TEXT, "space_after": 8}),
            ("Depois do Grokking: o personagem já está andando — a tecla virou reflexo.", {"size": 15, "color": TEXT, "space_after": 8}),
            ("Nesta aula, o Input Map + Play começam esse caminho.", {"size": 15, "color": TEXT_DIM}),
        ],
    )
    add_footer(s, W, H)

    # 6 — Assets
    s = new_slide(prs)
    add_num_badge(s, "05")
    add_title(s, "Assets — imagens e sons no FileSystem")
    items = [
        ("IMAGENS", "Sprites, ícones, tiles — o Sprite2D consome uma textura."),
        ("SONS", "SFX e música — recursos referenciados por nós de áudio."),
        ("OUTROS", "Fontes, cenas, scripts — tudo que o jogo consome no projeto."),
    ]
    for i, (k, v) in enumerate(items):
        top = Inches(1.15 + i * 1.15)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(1.05))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.2), Inches(8.4), Inches(0.7))
        write_lines(
            tf,
            [
                (k, {"size": 13, "bold": True, "color": GOLD, "font": TITLE_FONT, "space_after": 4}),
                (v, {"size": 15, "color": TEXT}),
            ],
        )
    add_footer(s, W, H)

    # 7 — Ponte
    s = new_slide(prs)
    add_num_badge(s, "06")
    add_title(s, "Ponte: do glossário à prática")
    _, tf = add_textbox(s, Inches(0.7), Inches(1.5), Inches(8.6), Inches(2.5))
    write_lines(
        tf,
        [
            ("Glossário nomeia o ofício;", {"size": 24, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("a cena do Player é o primeiro lugar", {"size": 22, "bold": True, "color": TEXT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("onde esses termos viram estrutura.", {"size": 22, "bold": True, "color": TEXT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 18}),
            ("Próximo: Cenas, Nós e a hierarquia do Jogador.", {"size": 14, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 8 — Cenas e Nós
    s = new_slide(prs)
    add_num_badge(s, "07")
    add_title(s, "Godot — Cenas e Nós (“receita de bolo”)")
    cards = [
        ("NÓ", "Tudo na Godot é um nó: corpo, sprite, colisão, script."),
        ("CENA", "Arquivo .tscn = receita reutilizável com esses nós."),
        ("ÁRVORE", "Filhos herdam a transform da raiz — hierarquia importa."),
    ]
    for i, (k, v) in enumerate(cards):
        left = Inches(0.5 + i * 3.15)
        add_panel(s, left, Inches(1.25), Inches(3.0), Inches(2.7))
        _, tf = add_textbox(s, left + Inches(0.2), Inches(1.55), Inches(2.6), Inches(2.2))
        write_lines(
            tf,
            [
                (k, {"size": 16, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
                (v, {"size": 14, "color": TEXT, "align": PP_ALIGN.CENTER}),
            ],
        )
    add_footer(s, W, H)

    # 9 — Hierarquia Player
    s = new_slide(prs)
    add_num_badge(s, "08")
    add_title(s, "Hierarquia do Player")
    nodes = [
        ("CharacterBody2D", "Raiz — corpo controlável do herói"),
        ("Sprite2D", "Filho — imagem / asset visual"),
        ("CollisionShape2D", "Filho — forma de colisão"),
    ]
    for i, (k, v) in enumerate(nodes):
        top = Inches(1.2 + i * 1.1)
        indent = Inches(0.7 + (0 if i == 0 else 0.55))
        add_panel(s, indent, top, Inches(8.2 if i == 0 else 7.65), Inches(0.95))
        _, tf = add_textbox(s, indent + Inches(0.25), top + Inches(0.18), Inches(7.5), Inches(0.65))
        write_lines(
            tf,
            [
                (k, {"size": 16, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "space_after": 2}),
                (v, {"size": 13, "color": TEXT_DIM}),
            ],
        )
    add_footer(s, W, H)

    # 10 — Passo a passo cena
    s = new_slide(prs)
    add_num_badge(s, "09")
    add_title(s, "Passo a passo: criar a cena do zero")
    steps = [
        "Nova cena com raiz CharacterBody2D (renomeie para Player).",
        "Adicione Sprite2D e atribua uma textura (placeholder da Godot serve).",
        "Adicione CollisionShape2D com RectangleShape2D ou CapsuleShape2D.",
        "Salve como player.tscn.",
    ]
    for i, text in enumerate(steps, start=1):
        top = Inches(1.1 + (i - 1) * 0.85)
        badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.55), top + Inches(0.1), Inches(0.5), Inches(0.5))
        badge.fill.solid()
        badge.fill.fore_color.rgb = BLOOD
        badge.line.color.rgb = GOLD
        tf = badge.text_frame
        p = _clear_paragraphs(tf)
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = f"{i:02d}"
        _set_run(run, size_pt=12, bold=True, color=GOLD_BRIGHT, font=TITLE_FONT)
        add_panel(s, Inches(1.25), top, Inches(8.1), Inches(0.7))
        _, tf = add_textbox(s, Inches(1.45), top + Inches(0.15), Inches(7.7), Inches(0.45))
        write_lines(tf, [text], size_pt=14, color=TEXT)
    add_footer(s, W, H)

    # 11 — Input Map ações
    s = new_slide(prs)
    add_num_badge(s, "10")
    add_title(s, "Input Map — ações ir_*")
    _, tf = add_textbox(s, Inches(0.6), Inches(1.05), Inches(8.8), Inches(0.45))
    write_lines(tf, ["Project → Project Settings → Input Map"], size_pt=14, color=TEXT_DIM, font=TITLE_FONT)
    actions = ["ir_cima", "ir_baixo", "ir_esquerda", "ir_direita"]
    for i, act in enumerate(actions):
        col = i % 2
        row = i // 2
        left = Inches(0.7 + col * 4.5)
        top = Inches(1.7 + row * 1.3)
        add_panel(s, left, top, Inches(4.2), Inches(1.1))
        _, tf = add_textbox(s, left + Inches(0.2), top + Inches(0.3), Inches(3.8), Inches(0.5))
        write_lines(tf, [(act, {"size": 20, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER})])
    add_footer(s, W, H)

    # 12 — Teclas
    s = new_slide(prs)
    add_num_badge(s, "11")
    add_title(s, "Teclas: WASD ou setas")
    add_panel(s, Inches(0.55), Inches(1.3), Inches(4.3), Inches(3.0))
    _, tf = add_textbox(s, Inches(0.8), Inches(1.6), Inches(3.8), Inches(2.5))
    write_lines(
        tf,
        [
            ("WASD", {"size": 20, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("W → ir_cima", {"size": 15, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("A → ir_esquerda", {"size": 15, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("S → ir_baixo", {"size": 15, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("D → ir_direita", {"size": 15, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_panel(s, Inches(5.15), Inches(1.3), Inches(4.3), Inches(3.0))
    _, tf = add_textbox(s, Inches(5.4), Inches(1.6), Inches(3.8), Inches(2.5))
    write_lines(
        tf,
        [
            ("SETAS", {"size": 20, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 12}),
            ("↑ → ir_cima", {"size": 15, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("← → ir_esquerda", {"size": 15, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("↓ → ir_baixo", {"size": 15, "align": PP_ALIGN.CENTER, "space_after": 6}),
            ("→ → ir_direita", {"size": 15, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    # 13 — Script
    s = new_slide(prs)
    add_num_badge(s, "12")
    add_title(s, "Script mínimo — player.gd + move_and_slide")
    add_panel(s, Inches(0.55), Inches(1.15), Inches(8.9), Inches(3.3))
    _, tf = add_textbox(s, Inches(0.85), Inches(1.4), Inches(8.3), Inches(2.9))
    write_lines(
        tf,
        [
            ("Anexe player.gd ao CharacterBody2D.", {"size": 16, "color": TEXT, "space_after": 10}),
            ("Leia as ações ir_* (Input.get_vector ou is_action_pressed).", {"size": 16, "color": TEXT, "space_after": 10}),
            ("Aplique velocity e chame move_and_slide().", {"size": 16, "color": TEXT, "space_after": 10}),
            ("Pressione Play — sinta o início do Grokking.", {"size": 16, "bold": True, "color": GOLD_BRIGHT, "space_after": 10}),
            ("Fora de escopo: câmera follow, jump avançado, animação.", {"size": 13, "color": TEXT_DIM}),
        ],
    )
    add_footer(s, W, H)

    # 14 — Checklist
    s = new_slide(prs)
    add_num_badge(s, "13")
    add_title(s, "Checklist do artefato")
    checks = [
        "Cena player.tscn salva com hierarquia correta.",
        "Sprite2D com textura e CollisionShape2D com shape.",
        "Input Map com as quatro ações ir_* mapeadas.",
        "player.gd move o personagem nas quatro direções no Play.",
        "Anotações + síntese enviadas na página da aula.",
    ]
    for i, text in enumerate(checks):
        top = Inches(1.15 + i * 0.65)
        add_panel(s, Inches(0.55), top, Inches(8.9), Inches(0.55))
        _, tf = add_textbox(s, Inches(0.8), top + Inches(0.1), Inches(8.4), Inches(0.4))
        write_lines(tf, [(f"□  {text}", {"size": 14, "color": TEXT})])
    add_footer(s, W, H)

    # 15 — Fechamento
    s = new_slide(prs)
    add_num_badge(s, "14")
    add_title(s, "Fechamento e próxima aula")
    add_panel(s, Inches(0.55), Inches(1.3), Inches(8.9), Inches(3.0))
    _, tf = add_textbox(s, Inches(0.9), Inches(1.7), Inches(8.2), Inches(2.4))
    write_lines(
        tf,
        [
            ("FIM DA AULA 02: GLOSSÁRIO E PLAYER", {"size": 18, "bold": True, "color": GOLD_BRIGHT, "font": TITLE_FONT, "align": PP_ALIGN.CENTER, "space_after": 16}),
            ("Você nomeou o ofício e colocou o herói na tela.", {"size": 15, "color": TEXT, "align": PP_ALIGN.CENTER, "space_after": 10}),
            ("Próxima aula (ponte): câmera follow e polish de movimento.", {"size": 15, "color": TEXT_DIM, "align": PP_ALIGN.CENTER, "space_after": 16}),
            ("Leve o código da aula ao Altar quando o Mestre liberar.", {"size": 13, "color": TEXT_DIM, "align": PP_ALIGN.CENTER}),
        ],
    )
    add_footer(s, W, H)

    return prs


def export_pdf_with_powerpoint(pptx_path: Path, pdf_path: Path) -> bool:
    """Exporta PDF via Microsoft PowerPoint (Windows COM)."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        # Fallback: PowerShell COM sem pywin32
        ps = f"""
$ppt = New-Object -ComObject PowerPoint.Application
$ppt.Visible = [Microsoft.Office.Core.MsoTriState]::msoTrue
$pres = $ppt.Presentations.Open('{pptx_path}', $true, $false, $false)
$pres.SaveAs('{pdf_path}', 32)
$pres.Close()
$ppt.Quit()
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($pres) | Out-Null
[System.Runtime.Interopservices.Marshal]::ReleaseComObject($ppt) | Out-Null
"""
        import subprocess

        result = subprocess.run(
            ["powershell", "-NoProfile", "-Command", ps],
            capture_output=True,
            text=True,
            check=False,
        )
        if result.returncode != 0:
            print(result.stdout)
            print(result.stderr, file=sys.stderr)
            return False
        return pdf_path.exists()

    powerpoint = win32com.client.Dispatch("PowerPoint.Application")
    powerpoint.Visible = 1
    presentation = powerpoint.Presentations.Open(str(pptx_path), WithWindow=False)
    # 32 = ppSaveAsPDF
    presentation.SaveAs(str(pdf_path), 32)
    presentation.Close()
    powerpoint.Quit()
    return pdf_path.exists()


def main():
    prs = build_deck()
    OUT_PPTX.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(OUT_PPTX))
    print(f"PPTX: {OUT_PPTX} ({OUT_PPTX.stat().st_size} bytes)")

    if OUT_PDF.exists():
        OUT_PDF.unlink()
    ok = export_pdf_with_powerpoint(OUT_PPTX.resolve(), OUT_PDF.resolve())
    if ok:
        print(f"PDF:  {OUT_PDF} ({OUT_PDF.stat().st_size} bytes)")
    else:
        print("AVISO: falha ao exportar PDF via PowerPoint.", file=sys.stderr)
        sys.exit(2)

    # Sanity: slide count
    check = Presentation(str(OUT_PPTX))
    print(f"Slides: {len(check.slides)}")


if __name__ == "__main__":
    main()
