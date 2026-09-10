#!/usr/bin/env python3
"""Gera o material Word da oficina do Player (Aula 02).

Saída:
  assets/docs/aulas/aula02-player/Material-Player-Aula02.docx
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import nsmap, qn
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "docs" / "aulas" / "aula02-player" / "Material-Player-Aula02.docx"

GOLD = RGBColor(0x7A, 0x5C, 0x2E)
TEXT = RGBColor(0x1A, 0x14, 0x12)
DIM = RGBColor(0x4A, 0x3F, 0x38)

# Namespace dos content controls de checkbox (Word 2010+)
W14_NS = "http://schemas.microsoft.com/office/word/2010/wordml"
if "w14" not in nsmap:
    nsmap["w14"] = W14_NS


def _set_run_font(run, *, size_pt=11, bold=False, italic=False, color=TEXT, name="Calibri"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        from docx.oxml import OxmlElement

        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def _add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        _set_run_font(run, size_pt=18 if level == 1 else 14, bold=True, color=GOLD, name="Georgia")
    return p


def _add_para(doc: Document, text: str, *, italic=False, size=11):
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size_pt=size, italic=italic, color=TEXT)
    p.paragraph_format.space_after = Pt(8)
    return p


def _add_bullet(doc: Document, text: str):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    _set_run_font(run, size_pt=11)
    return p


def _checkbox_sdt(*, checked: bool = False) -> OxmlElement:
    """Content control de checkbox clicável no Word (Desktop)."""
    sdt = OxmlElement("w:sdt")
    sdt_pr = OxmlElement("w:sdtPr")

    id_el = OxmlElement("w:id")
    id_el.set(qn("w:val"), str(abs(hash((checked, id(sdt)))) % 2_000_000_000))
    sdt_pr.append(id_el)

    checkbox = OxmlElement("w14:checkbox")
    checked_el = OxmlElement("w14:checked")
    checked_el.set(qn("w14:val"), "1" if checked else "0")

    checked_state = OxmlElement("w14:checkedState")
    checked_state.set(qn("w14:val"), "2612")  # ☒
    checked_state.set(qn("w14:font"), "Segoe UI Symbol")

    unchecked_state = OxmlElement("w14:uncheckedState")
    unchecked_state.set(qn("w14:val"), "2610")  # ☐
    unchecked_state.set(qn("w14:font"), "Segoe UI Symbol")

    checkbox.append(checked_el)
    checkbox.append(checked_state)
    checkbox.append(unchecked_state)
    sdt_pr.append(checkbox)
    sdt.append(sdt_pr)

    sdt_content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Segoe UI Symbol")
    r_fonts.set(qn("w:hAnsi"), "Segoe UI Symbol")
    r_pr.append(r_fonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    r_pr.append(sz)
    run.append(r_pr)

    text_el = OxmlElement("w:t")
    text_el.text = "\u2612" if checked else "\u2610"
    run.append(text_el)
    sdt_content.append(run)
    sdt.append(sdt_content)
    return sdt


def _add_checkbox_item(doc: Document, text: str, *, checked: bool = False):
    """Linha com checkbox interativo + rótulo (clique na caixinha no Word)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p._p.append(_checkbox_sdt(checked=checked))

    spacer = p.add_run("  ")
    _set_run_font(spacer, size_pt=11)
    label = p.add_run(text)
    _set_run_font(label, size_pt=11)
    return p


def _add_numbered(doc: Document, text: str):
    p = doc.add_paragraph(style="List Number")
    run = p.add_run(text)
    _set_run_font(run, size_pt=11)
    return p


def _add_code_block(doc: Document, code: str):
    for line in code.splitlines() or [""]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(12)
        run = p.add_run(line if line else " ")
        _set_run_font(run, size_pt=10, name="Consolas", color=TEXT)
    doc.add_paragraph()


def build() -> Path:
    doc = Document()

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("Aula 02 — Material do Player (Godot 4)")
    _set_run_font(run, size_pt=22, bold=True, color=GOLD, name="Georgia")

    _add_para(
        doc,
        "Passo a passo para montar a primeira cena do Jogador: hierarquia de nós, "
        "Input Map e script mínimo com move_and_slide.",
    )
    _add_para(
        doc,
        "Não há ZIP de sprite nesta aula — use um ícone/placeholder da própria Godot no Sprite2D.",
        italic=True,
    )

    _add_heading(doc, "Checklist do artefato", level=2)
    _add_para(doc, "Ao final, você deve ter (clique nas caixinhas para marcar):")
    for item in [
        "Cena salva (ex.: player.tscn)",
        "Raiz CharacterBody2D (renomeada para Player)",
        "Filho Sprite2D com alguma textura",
        "Filho CollisionShape2D com shape alinhado ao sprite",
        "Input Map com ir_cima, ir_baixo, ir_esquerda, ir_direita",
        "Teclas WASD ou setas ligadas às ações",
        "Script player.gd anexado à raiz e movimento funcionando no Play",
    ]:
        _add_checkbox_item(doc, item)
    _add_para(
        doc,
        "Registre tudo nas anotações da página da aula (glossário + hierarquia + Input Map + observações).",
        italic=True,
    )

    _add_heading(doc, "1. Cenas e nós", level=2)
    _add_para(
        doc,
        "Na Godot, tudo é um nó. Uma cena (.tscn) é a “receita de bolo” que guarda essa hierarquia para reutilizar.",
    )
    _add_para(doc, "Painéis úteis:", size=11)
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    headers = ("Painel", "Função")
    rows = [
        ("Scene", "Hierarquia dos nós"),
        ("FileSystem", "Assets (imagens, sons, scripts, cenas)"),
        ("Inspector", "Propriedades do nó selecionado"),
        ("Viewport 2D", "Visualização da cena"),
    ]
    for i, cell_text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(cell_text)
        _set_run_font(run, bold=True, size_pt=11, color=GOLD)
    for r, (a, b) in enumerate(rows, start=1):
        table.rows[r].cells[0].text = ""
        table.rows[r].cells[1].text = ""
        ra = table.rows[r].cells[0].paragraphs[0].add_run(a)
        rb = table.rows[r].cells[1].paragraphs[0].add_run(b)
        _set_run_font(ra, size_pt=11)
        _set_run_font(rb, size_pt=11)
    doc.add_paragraph()

    _add_heading(doc, "2. Criar a cena do Player", level=2)
    steps = [
        "Scene → New Scene → escolha CharacterBody2D como raiz.",
        "Renomeie a raiz para Player.",
        "Clique com o botão direito na raiz → Add Child Node: Sprite2D e CollisionShape2D.",
        "No Sprite2D, em Texture, use um placeholder (ex.: ícone embutido da Godot ou uma imagem simples do projeto).",
        "No CollisionShape2D, em Shape, crie um RectangleShape2D ou CapsuleShape2D e ajuste ao tamanho do sprite.",
        "Salve a cena como player.tscn (ou nome combinado com a turma).",
    ]
    for step in steps:
        _add_numbered(doc, step)

    _add_para(doc, "Hierarquia esperada:")
    _add_code_block(
        doc,
        "Player (CharacterBody2D)\n├── Sprite2D\n└── CollisionShape2D",
    )

    _add_heading(doc, "3. Input Map", level=2)
    for step in [
        "Project → Project Settings → Input Map.",
        "Crie as ações (digite o nome e clique Add): ir_cima, ir_baixo, ir_esquerda, ir_direita.",
        "Em cada ação, clique + e associe WASD (W / S / A / D) ou setas (↑ / ↓ / ← / →).",
    ]:
        _add_numbered(doc, step)
    _add_para(
        doc,
        "Use nomes internos (ir_*) no código — assim trocar teclas depois não quebra o script.",
        italic=True,
    )

    _add_heading(doc, "4. Script mínimo (player.gd)", level=2)
    for step in [
        "Selecione o CharacterBody2D (Player).",
        "Clique em Attach Script e salve como player.gd.",
        "Substitua o conteúdo por algo equivalente ao código abaixo.",
        "Pressione Play (F5). Se a Godot pedir uma cena principal, escolha player.tscn (ou uma cena de teste que instancia o Player).",
        "Confirme movimento nas quatro direções — esse é o começo do Grokking.",
    ]:
        _add_numbered(doc, step)

    _add_para(doc, "Código sugerido:")
    _add_code_block(
        doc,
        """extends CharacterBody2D

@export var speed: float = 200.0

func _physics_process(_delta: float) -> void:
	var direction := Input.get_vector("ir_esquerda", "ir_direita", "ir_cima", "ir_baixo")
	velocity = direction * speed
	move_and_slide()""",
    )

    _add_para(doc, "Fora de escopo nesta aula:", italic=True)
    for item in ["Câmera follow", "Pulo avançado / gravidade de plataforma", "Animações de sprite"]:
        _add_bullet(doc, item)

    _add_heading(doc, "5. Ligação com o glossário", level=2)
    gtable = doc.add_table(rows=4, cols=2)
    gtable.style = "Table Grid"
    gheaders = ("Termo", "Onde aparece na prática")
    grows = [
        ("Core Loop", "Andar na tela é o primeiro passo do ciclo (depois virão coletar / avançar)."),
        ("Grokking", "Quando ir_* deixa de ser “pensar a tecla” e vira reflexo no Play."),
        ("Assets", "A textura do Sprite2D (e futuros sons) no FileSystem."),
    ]
    for i, cell_text in enumerate(gheaders):
        cell = gtable.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(cell_text)
        _set_run_font(run, bold=True, size_pt=11, color=GOLD)
    for r, (a, b) in enumerate(grows, start=1):
        gtable.rows[r].cells[0].text = ""
        gtable.rows[r].cells[1].text = ""
        ra = gtable.rows[r].cells[0].paragraphs[0].add_run(a)
        rb = gtable.rows[r].cells[1].paragraphs[0].add_run(b)
        _set_run_font(ra, size_pt=11, bold=True)
        _set_run_font(rb, size_pt=11)
    doc.add_paragraph()

    _add_heading(doc, "Dúvidas comuns", level=2)

    p = doc.add_paragraph()
    run = p.add_run("O personagem não se move")
    _set_run_font(run, bold=True, size_pt=11)
    _add_para(
        doc,
        "Confira os nomes das ações no Input Map (iguais ao script) e se o script está anexado à raiz CharacterBody2D.",
    )

    p = doc.add_paragraph()
    run = p.add_run("Não colide com nada")
    _set_run_font(run, bold=True, size_pt=11)
    _add_para(
        doc,
        "Nesta aula o foco é estrutura + input. Sem chão/StaticBody2D, o Player ainda deve deslizar no vazio do Viewport.",
    )

    p = doc.add_paragraph()
    run = p.add_run("Posso usar outro shape?")
    _set_run_font(run, bold=True, size_pt=11)
    _add_para(doc, "Sim — o importante é existir um CollisionShape2D com shape definido.")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"OK: {path}")
